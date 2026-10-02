"""tests/test_pipeline.py

End-to-end tests for the Word Citation Footnote Agent. Every test builds a
small .docx, runs the real pipeline (main.main), and inspects the actual
OOXML (footnotes.xml, document.xml, rels) - not just visible python-docx text.

Run:  python -m pytest tests/ -v
"""
import hashlib
import re
import sys
import zipfile
from pathlib import Path

import pytest
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document                                    # noqa: E402
from docx.oxml import OxmlElement                            # noqa: E402
from docx.oxml.ns import qn                                  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT       # noqa: E402
from docx.opc.packuri import PackURI                         # noqa: E402
from docx.opc.part import Part                               # noqa: E402

import main as agent_main                                    # noqa: E402
from document_reader import (iter_paragraphs, raw_text_of,  # noqa: E402
                             snapshot_existing_footnotes)
from reference_parser import ReferenceParser                 # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

REFS = [
    "1. Alpha. One. https://example.com/one",
    "2. Beta. Two. https://example.com/two",
    "3. Gamma. Three. https://example.com/three",
    "4. Delta. Four. https://example.com/four",
]


# --------------------------------------------------------------------------
# Builders
# --------------------------------------------------------------------------
def _add_runs(p, parts):
    for part in parts:
        if isinstance(part, str):
            p.add_run(part)
        else:  # (text, bold)
            r = p.add_run(part[0])
            r.bold = part[1]


def make_doc(path, body, refs=REFS, heading="References",
             heading_styled=True, extra=None):
    """body: list of items; str = plain paragraph; list = runs (str or
    (text, bold)). ``extra`` is an optional callback(doc) for custom XML."""
    doc = Document()
    for item in body:
        if isinstance(item, str):
            doc.add_paragraph(item)
        else:
            p = doc.add_paragraph()
            _add_runs(p, item)
    if extra:
        extra(doc)  # custom structure belongs in the BODY, before the refs
    if heading is not None:
        if heading_styled:
            doc.add_heading(heading, level=1)
        else:
            doc.add_paragraph(heading)
    for r in (refs or []):
        doc.add_paragraph(r)
    doc.save(str(path))
    return Path(path)


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph):
    r1 = OxmlElement("w:r")
    fc1 = OxmlElement("w:fldChar")
    fc1.set(qn("w:fldCharType"), "begin")
    r1.append(fc1)
    r2 = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.text = " PAGE "
    r2.append(it)
    r3 = OxmlElement("w:r")
    fc2 = OxmlElement("w:fldChar")
    fc2.set(qn("w:fldCharType"), "end")
    r3.append(fc2)
    for r in (r1, r2, r3):
        paragraph._p.append(r)


def add_bookmark(paragraph, bid="1", name="bm"):
    bs = OxmlElement("w:bookmarkStart")
    bs.set(qn("w:id"), bid)
    bs.set(qn("w:name"), name)
    be = OxmlElement("w:bookmarkEnd")
    be.set(qn("w:id"), bid)
    paragraph._p.append(bs)
    paragraph._p.append(be)


FOOTNOTES_CT = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml")
FOOTNOTES_TEMPLATE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<w:footnotes '
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="0">'
    '<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    "</w:footnotes>")


def add_preexisting_footnote(doc, paragraph, text, fid="1"):
    part = Part(PackURI("/word/footnotes.xml"), FOOTNOTES_CT,
                FOOTNOTES_TEMPLATE.encode("utf-8"), doc.part.package)
    doc.part.relate_to(part, RT.FOOTNOTES)
    root = etree.fromstring(FOOTNOTES_TEMPLATE.encode("utf-8"))
    ft = OxmlElement("w:footnote")
    ft.set(qn("w:id"), fid)
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:footnoteRef"))
    p.append(r)
    r2 = OxmlElement("w:r")
    t2 = OxmlElement("w:t")
    t2.text = text
    r2.append(t2)
    p.append(r2)
    ft.append(p)
    root.append(ft)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                standalone=True)
    run = OxmlElement("w:r")
    fr = OxmlElement("w:footnoteReference")
    fr.set(qn("w:id"), fid)
    run.append(fr)
    paragraph._p.append(run)


# --------------------------------------------------------------------------
# Pipeline helpers
# --------------------------------------------------------------------------
def run_tool(input_path, tmp_path, *extra_args):
    out_dir = tmp_path / "out"
    rc = agent_main.main(["--input", str(input_path),
                          "--output-dir", str(out_dir), *extra_args])
    outputs = sorted(out_dir.glob("*.docx")) if out_dir.exists() else []
    return rc, outputs


def output_path(tmp_path):
    out_dir = tmp_path / "out"
    docs = sorted(out_dir.glob("*.docx"))
    assert docs, "no output produced"
    return docs[0]


def load(path):
    doc = Document(path)
    z = zipfile.ZipFile(path)
    ft = None
    for n in z.namelist():
        if n.endswith("footnotes.xml"):
            ft = etree.fromstring(z.read(n))
    return doc, z, ft


def regular(ft):
    if ft is None:
        return []
    return [f for f in ft.findall(qn("w:footnote"))
            if f.get(qn("w:id")) not in ("-1", "0")]


def fn_text(f):
    return "".join(t.text or "" for t in f.iter(qn("w:t")))


def body_ref_ids(doc):
    return [r.get(qn("w:id")) for r in doc.element.body.iter(qn("w:footnoteReference"))]


def body_text_before_refs(doc):
    """Raw text of all paragraphs up to (excluding) the references heading."""
    parts = []
    for p, _w in _iter(doc):
        t = raw_text_of(p._p)
        if re.match(r"^\s*(\d+\.?\s+)?(References|Bibliography|Works Cited|Sources)\s*$",
                    t, re.I):
            break
        parts.append(t)
    return "\n".join(parts)


def _iter(doc):
    from document_reader import iter_paragraph_info
    return iter_paragraph_info(doc)


def read_csv(tmp_path):
    import csv
    p = tmp_path / "out" / "citation_review_report.csv"
    assert p.exists(), "report CSV missing"
    with open(p, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    summary = {}
    in_sum = False
    for row in rows:
        if row and row[0] == "SUMMARY":
            in_sum = True
            continue
        if in_sum and len(row) >= 2:
            summary[row[0]] = row[1]
    return rows, summary


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def no_markers(text):
    return not re.search(r"\[\d+\]", text)


# --------------------------------------------------------------------------
# 1-5: basic citation syntax
# --------------------------------------------------------------------------
def test_simple_citation(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["See [1] for details."])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    assert fn_text(regular(ft)[0]) == "https://example.com/one"
    assert len(body_ref_ids(doc)) == 1
    assert no_markers(body_text_before_refs(doc))


def test_multiple_citations(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["A [1] and [2] B."])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    urls = {fn_text(f) for f in regular(ft)}
    assert urls == {"https://example.com/one", "https://example.com/two"}
    assert len(body_ref_ids(doc)) == 2
    assert no_markers(body_text_before_refs(doc))


def test_repeated_citation_numbers(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["First [2].", "Again [2]."])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 2
    assert all(fn_text(f) == "https://example.com/two" for f in regular(ft))
    assert len(body_ref_ids(doc)) == 2


def test_range_hyphen(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["Range [1]-[3] here."])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    urls = [fn_text(f) for f in sorted(regular(ft), key=lambda f: int(f.get(qn("w:id"))))]
    assert urls == ["https://example.com/one",
                    "https://example.com/two",
                    "https://example.com/three"]
    assert len(body_ref_ids(doc)) == 3
    assert no_markers(body_text_before_refs(doc))


def test_range_en_dash(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["Range [1]\u2013[3] and [1] \u2013 [3] too."])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 6
    assert no_markers(body_text_before_refs(doc))


def test_range_em_dash(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["Range [1]\u2014[3] here."])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 3


def test_reversed_range_is_ambiguous(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["Odd [3]-[1] range."])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 0
    assert "[3]-[1]" in body_text_before_refs(doc)  # left unchanged
    _rows, summary = read_csv(tmp_path)
    assert summary["Ambiguous_citations"] == "1"


# --------------------------------------------------------------------------
# 6-8: formatting + tables + hyperlinks
# --------------------------------------------------------------------------
def test_citation_in_bold_text(tmp_path):
    src = make_doc(tmp_path / "t.docx",
                   [["Plain ", ("Bold [2] claim", True), " plain."]])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    assert fn_text(regular(ft)[0]) == "https://example.com/two"
    # bold must survive on both sides of the footnote reference
    bold_texts = "".join(r.text for p in doc.paragraphs for r in p.runs if r.bold)
    assert "Bold " in bold_texts and " claim" in bold_texts
    # the footnote reference run itself keeps the bold formatting
    z = zipfile.ZipFile(output_path(tmp_path))
    xml = z.read("word/document.xml").decode()
    ref_run = re.search(r'<w:r>(?:(?!</w:r>).)*?footnoteReference(?:(?!</w:r>).)*?</w:r>',
                        xml, re.S)
    assert ref_run and "<w:b/>" in ref_run.group(0)
    assert no_markers(body_text_before_refs(doc))


def test_citation_in_table(tmp_path):
    def build(doc):
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "Header"
        tbl.cell(1, 1).text = "Cell cites [3]."
    src = make_doc(tmp_path / "t.docx", ["Intro text."], extra=build)
    rc, _ = run_tool(tmp_path / "t.docx", tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    assert fn_text(regular(ft)[0]) == "https://example.com/three"
    tbl = doc.tables[0]
    assert "Cell cites" in tbl.cell(1, 1).text
    assert no_markers(tbl.cell(1, 1).text)
    assert len(doc.tables) == 1  # table preserved


def test_hyperlink_preserved(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Before link [1]. ")
    add_hyperlink(p, "https://example.com/external", " (ext)")
    p.add_run(" After [1].")
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = Path(tmp_path / "t.docx")
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc2, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 2  # both [1] citations (outside the link)
    hls = doc2.element.body.findall(".//" + qn("w:hyperlink"))
    assert len(hls) == 1
    targets = [doc2.part.rels[h.get(qn("r:id"))].target_ref for h in hls]
    assert targets == ["https://example.com/external"]
    assert "(ext)" in body_text_before_refs(doc2)


# --------------------------------------------------------------------------
# 9-11: missing data must be reported, never invented
# --------------------------------------------------------------------------
def test_missing_reference_url(tmp_path):
    refs = ["1. Alpha. One. https://example.com/one",
            "2. Beta. Two. no url here."]
    src = make_doc(tmp_path / "t.docx", ["Cite [1] and [2]."], refs=refs)
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1  # only [1]
    assert fn_text(regular(ft)[0]) == "https://example.com/one"
    body = body_text_before_refs(doc)
    assert "[2]" in body  # left unchanged
    assert "[1]" not in body
    _rows, summary = read_csv(tmp_path)
    assert summary["Unresolved_citations"] == "1"
    assert summary["Missing_URLs"] == "2"


def test_missing_reference_number(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["Ghost [99] and real [1]."])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    body = body_text_before_refs(doc)
    assert "[99]" in body  # left unchanged
    _rows, summary = read_csv(tmp_path)
    assert summary["Unresolved_citations"] == "1"


def test_uncited_reference_reported(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["Only [4] cited."])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _rows, summary = read_csv(tmp_path)
    assert summary["Uncited_references"] == "1, 2, 3"


def test_reference_section_not_scanned_for_citations(tmp_path):
    # reference list itself uses [n] numbering; it must NOT be treated as body
    refs = ["[1] Alpha. One. https://example.com/one",
            "[2] Beta. Two. https://example.com/two"]
    src = make_doc(tmp_path / "t.docx", ["Body cites [1] only."], refs=refs)
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    assert fn_text(regular(ft)[0]) == "https://example.com/one"


# --------------------------------------------------------------------------
# 12-13: safety (hyperlink-internal + cross-run markers)
# --------------------------------------------------------------------------
def test_citation_inside_hyperlink_untouched(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Link with cite: ")
    add_hyperlink(p, "https://example.com/page", "see [2] here")
    p.add_run(" end. And real [1].")
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = Path(tmp_path / "t.docx")
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc2, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1  # only the safe [1]
    hls = doc2.element.body.findall(".//" + qn("w:hyperlink"))
    hl_text = "".join(t.text or "" for t in hls[0].iter(qn("w:t")))
    assert hl_text == "see [2] here"  # hyperlink content untouched
    _rows, summary = read_csv(tmp_path)
    assert summary["Ambiguous_citations"] == "1"


def test_citation_split_across_runs_untouched(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("See [1")
    p.add_run("] for detail. Safe [2].")
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = Path(tmp_path / "t.docx")
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc2, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1  # only the safe [2]
    urls = {fn_text(f) for f in regular(ft)}
    assert urls == {"https://example.com/two"}
    # the split marker must remain in two separate runs, untouched
    body = body_text_before_refs(doc2)
    assert "See [1" in body and "] for detail" in body
    _rows, summary = read_csv(tmp_path)
    assert summary["Ambiguous_citations"] == "1"


def test_field_and_bookmark_preserved(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Text [1] with ")
    add_field(p)
    p.add_run(" field and ")
    add_bookmark(p)
    p.add_run(" more.")
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = Path(tmp_path / "t.docx")
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc2, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    xml = _z.read("word/document.xml").decode()
    assert xml.count("<w:fldChar") == 2  # begin + end preserved
    assert "<w:instrText" in xml and "PAGE" in xml  # field instruction intact
    assert xml.count("<w:bookmarkStart") == 1
    assert xml.count("<w:bookmarkEnd") == 1
    assert no_markers(body_text_before_refs(doc2))


# --------------------------------------------------------------------------
# 14-15: existing footnotes + idempotency
# --------------------------------------------------------------------------
def test_existing_footnotes_preserved(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Intro with a manual note.")
    add_preexisting_footnote(doc, p, "Pre-existing footnote text.")
    doc.add_paragraph("And a citation [1].")
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = Path(tmp_path / "t.docx")
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc2, _z, ft = load(output_path(tmp_path))
    defs = {f.get(qn("w:id")): fn_text(f) for f in regular(ft)}
    assert "Pre-existing footnote text." in defs.values()
    assert "https://example.com/one" in defs.values()
    assert len(defs) == 2
    ids = body_ref_ids(doc2)
    assert len(ids) == 2
    _rows, summary = read_csv(tmp_path)
    assert summary["Existing_footnotes_detected"] == "1"
    assert summary["Existing_footnotes_preserved"] == "True"


def test_unrelated_existing_footnote_does_not_block(tmp_path):
    # A pre-existing footnote sits right before the citation's run; its
    # content does NOT match the citation's URL, so the citation must
    # still be processed (no false "previously processed" skip).
    doc = Document()
    p = doc.add_paragraph()
    add_preexisting_footnote(doc, p, "Unrelated note.")
    p.add_run(" [1] after note.")
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = Path(tmp_path / "t.docx")
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc2, _z, ft = load(output_path(tmp_path))
    defs = {fn_text(f) for f in regular(ft)}
    assert defs == {"Unrelated note.", "https://example.com/one"}
    _rows, summary = read_csv(tmp_path)
    assert summary["Previously_processed"] == "0"


def test_idempotent_second_run_replace(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["A [1] and [2] B."])
    rc1, _ = run_tool(src, tmp_path)
    assert rc1 == 0
    first = output_path(tmp_path)
    (tmp_path / "second.docx").write_bytes(first.read_bytes())
    rc2, out2 = run_tool(tmp_path / "second.docx", tmp_path / "r2")
    assert rc2 == 0
    _d, _z, ft2 = load(out2[0])
    assert len(regular(ft2)) == 2  # no duplicates
    _rows, summary = read_csv(tmp_path / "r2")
    assert summary["Footnotes_inserted"] == "0"
    assert summary["Previously_processed"] == "0"


def test_idempotent_second_run_keep_marker(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["A [1] and [2] B."])
    rc1, _ = run_tool(src, tmp_path, "--keep-marker")
    assert rc1 == 0
    first = output_path(tmp_path)
    r2 = tmp_path / "r2"
    rc2, out2 = run_tool(first, r2, "--keep-marker")
    assert rc2 == 0
    _d, _z, ft2 = load(out2[0])
    assert len(regular(ft2)) == 2  # no duplicates
    doc2, _z, _ft = load(out2[0])
    assert "[1]" in body_text_before_refs(doc2)  # marker kept
    _rows, summary = read_csv(r2)
    assert summary["Footnotes_inserted"] == "0"
    assert summary["Previously_processed"] == "2"


def test_input_hash_unchanged(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["A [1] B."])
    before = sha256(src)
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    assert sha256(src) == before


# --------------------------------------------------------------------------
# 16-18: reference-section detection
# --------------------------------------------------------------------------
def test_nonstandard_heading_bibliography(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["A [1] B."], heading="Bibliography")
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _d, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    _rows, summary = read_csv(tmp_path)
    assert summary["Reference_section_confidence"] == "high"


def test_numbered_heading_5_references(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["A [1] B."], heading="5. References")
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _d, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    _rows, summary = read_csv(tmp_path)
    assert summary["Reference_section_confidence"] == "high"


def test_bold_unstyled_heading(tmp_path):
    def build(doc):
        p = doc.add_paragraph("Works Cited")
        p.runs[0].bold = True
    doc = Document()
    doc.add_paragraph("A [2] B.")
    p = doc.add_paragraph("Works Cited")
    p.runs[0].bold = True
    for r in REFS:
        doc.add_paragraph(r)
    src = Path(tmp_path / "t.docx")
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _d, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1


def test_block_fallback_without_heading(tmp_path):
    # No heading at all; numbered block near the end.
    body = ["Prose line one.",
            "Prose line two cites [1].",
            "Prose line three.",
            "1. Alpha. One. https://example.com/one",
            "2. Beta. Two. https://example.com/two",
            "3. Gamma. Three. https://example.com/three"]
    src = make_doc(tmp_path / "t.docx", body, refs=None, heading=None)
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    assert fn_text(regular(ft)[0]) == "https://example.com/one"
    assert "Prose line two cites" in body_text_before_refs(doc)
    _rows, summary = read_csv(tmp_path)
    assert summary["Reference_section_method"] == "block-detection"


def test_no_reference_section_aborts(tmp_path):
    src = make_doc(tmp_path / "t.docx",
                   ["Just prose. Another line. A third line."],
                   refs=None, heading=None)
    before = sha256(src)
    rc, outputs = run_tool(src, tmp_path)
    assert rc == 1
    assert outputs == []  # no partial output
    assert sha256(src) == before
    _rows, summary = read_csv(tmp_path)
    assert "ABORTED" in summary["Status"]
    assert summary["Reference_section_method"] == "none"


def test_keyword_sentence_not_treated_as_heading(tmp_path):
    # The trap from the bug report: a short body sentence containing
    # "sources" must NOT be detected as the reference section; the real
    # numbered block below must be used instead.
    body = ["Our data sources were collected in 2019 from three labs.",
            "We cite [1] a lot here.",
            "1. Alpha. One. https://example.com/one",
            "2. Beta. Two. https://example.com/two",
            "3. Gamma. Three. https://example.com/three"]
    src = make_doc(tmp_path / "t.docx", body, refs=None, heading=None)
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    assert no_markers(body_text_before_refs(doc))
    _rows, summary = read_csv(tmp_path)
    assert summary["Reference_section_method"] == "block-detection"


def test_year_starting_line_not_an_entry(tmp_path):
    refs = ["1. Alpha. One. https://example.com/one",
            "2020. A journal article that starts with a year. https://year.example/x"]
    src = make_doc(tmp_path / "t.docx", ["A [1] B."], refs=refs)
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _d, _z, ft = load(output_path(tmp_path))
    urls = {fn_text(f) for f in regular(ft)}
    assert urls == {"https://example.com/one"}
    _rows, summary = read_csv(tmp_path)
    assert summary["References_detected"] == "1"


def test_url_selection_doi_preference(tmp_path):
    refs = ["1. Alpha. DOI https://doi.org/10.1/x plus https://example.com/direct",
            "2. Beta. Only a DOI url https://doi.org/10.2/y"]
    src = make_doc(tmp_path / "t.docx", ["A [1] B. C [2] D."], refs=refs)
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _d, _z, ft = load(output_path(tmp_path))
    urls = {fn_text(f) for f in regular(ft)}
    assert "https://example.com/direct" in urls       # direct preferred
    assert "https://doi.org/10.1/x" not in urls
    assert "https://doi.org/10.2/y" in urls           # explicit DOI URL used


def test_bare_doi_text_never_invented(tmp_path):
    refs = ["1. Alpha. DOI: 10.1000/abc. No http url at all."]
    src = make_doc(tmp_path / "t.docx", ["A [1] B."], refs=refs)
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _d, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 0
    body = body_text_before_refs(_d)
    assert "[1]" in body  # left unchanged, nothing invented


# --------------------------------------------------------------------------
# 19+: keep-marker, output safety, structures
# --------------------------------------------------------------------------
def test_keep_marker_mode(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["A [1] B."])
    rc, _ = run_tool(src, tmp_path, "--keep-marker")
    assert rc == 0
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 1
    body = body_text_before_refs(doc)
    assert "[1]" in body  # marker kept
    # the footnote reference run must come AFTER the marker text
    xml = _z.read("word/document.xml").decode()
    i_marker = xml.find(">[1]<")
    i_ref = xml.find("footnoteReference")
    assert i_marker != -1 and i_ref > i_marker


def test_output_collision_gets_safe_alternative(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["A [1] B."])
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    (out_dir / "t_with_footnotes.docx").write_bytes(b"pre-existing")
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    assert (out_dir / "t_with_footnotes_1.docx").exists()
    # and the pre-existing file was not clobbered
    assert (out_dir / "t_with_footnotes.docx").read_bytes() == b"pre-existing"


def test_overwrite_flag(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["A [1] B."])
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    (out_dir / "t_with_footnotes.docx").write_bytes(b"pre-existing")
    rc, _ = run_tool(src, tmp_path, "--overwrite")
    assert rc == 0
    assert (out_dir / "t_with_footnotes_1.docx").exists() is False
    assert zipfile.is_zipfile(out_dir / "t_with_footnotes.docx")


def test_output_path_equal_to_input_rejected(tmp_path):
    src = make_doc(tmp_path / "t.docx", ["A [1] B."])
    before = sha256(src)
    rc = agent_main.main(["--input", str(src), "--output", str(src)])
    assert rc == 2
    assert sha256(src) == before


def test_nested_table_citation(tmp_path):
    doc = Document()
    doc.add_paragraph("Outer [1].")
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Deep [2]."
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = Path(tmp_path / "t.docx")
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc2, _z, ft = load(output_path(tmp_path))
    urls = {fn_text(f) for f in regular(ft)}
    assert urls == {"https://example.com/one", "https://example.com/two"}


def test_text_fidelity_no_text_lost(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Before [1] ")
    add_hyperlink(p, "https://example.com/h", "link text")
    p.add_run(" and [2] after.")
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = Path(tmp_path / "t.docx")
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    from validator import direct_to_raw_map, ValidationContext
    from document_reader import iter_paragraph_info
    out = output_path(tmp_path)
    in_paras = [p._p for p, _w in iter_paragraph_info(Document(src))]
    out_paras = [p._p for p, _w in iter_paragraph_info(Document(out))]
    assert len(in_paras) == len(out_paras)
    for in_p, out_p in zip(in_paras, out_paras):
        expected = raw_text_of(in_p)
        # remove the two markers (they are SUCCESS in replace mode)
        for marker in ("[1]", "[2]"):
            if expected.count(marker):
                expected = expected.replace(marker, "", 1)
        assert raw_text_of(out_p) == expected, (raw_text_of(out_p), expected)


def test_com_backend_rejected_on_linux(tmp_path):
    if sys.platform == "win32":
        pytest.skip("COM backend only testable on Windows")
    src = make_doc(tmp_path / "t.docx", ["A [1] B."])
    before = sha256(src)
    rc, outputs = run_tool(src, tmp_path, "--backend", "com")
    assert rc == 2  # clear usage error, loud failure
    assert sha256(src) == before


def test_sample_end_to_end(tmp_path):
    """The sample document (freshly generated via tests/create_sample.py) must
    pass end-to-end: 11 new footnotes + 1 pre-existing preserved, input
    untouched. Generated into tmp_path - running the suite never modifies
    the tracked input/sample.docx."""
    import subprocess
    import sys as _sys
    venv = ROOT / ".venv"
    py = venv / ("Scripts/python.exe" if sys.platform == "win32" else "bin/python")
    if not py.exists():
        py = _sys.executable
    local = tmp_path / "sample.docx"
    tracked = ROOT / "input" / "sample.docx"
    tracked_before = sha256(tracked) if tracked.exists() else None
    proc = subprocess.run(
        [str(py), str(ROOT / "tests" / "create_sample.py"), str(local)],
        cwd=ROOT, capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr
    assert local.exists()
    if tracked_before is not None:
        assert sha256(tracked) == tracked_before  # repo file left alone
    before = sha256(local)
    rc, _ = run_tool(local, tmp_path)
    assert rc == 0
    assert sha256(local) == before
    doc, _z, ft = load(output_path(tmp_path))
    assert len(regular(ft)) == 12  # 11 new + 1 pre-existing
    _rows, summary = read_csv(tmp_path)
    assert summary["Footnotes_inserted"] == "11"
    assert summary["Existing_footnotes_preserved"] == "True"
    assert summary["Input_unchanged"] == "True"


# --------------------------------------------------------------------------
# Regressions found on real Word documents / realistic structures
# --------------------------------------------------------------------------
def _run_from_spec(spec):
    """A w:r built from ``spec``: every '|' is a w:lastRenderedPageBreak at
    that text position (Word's layout-cache hint), the rest is w:t text."""
    r = OxmlElement("w:r")
    for i, piece in enumerate(spec.split("|")):
        if i:
            r.append(OxmlElement("w:lastRenderedPageBreak"))
        if piece:
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = piece
            r.append(t)
    return r


URL_NUM = {r.split()[-1]: int(r.split(".")[0]) for r in REFS}   # url -> ref number


def _stream(p_elem, fn_url):
    """Linearised paragraph: characters, '|' for a page-break hint and '{n}'
    for a footnote reference whose footnote holds reference n's URL."""
    out = []
    for el in p_elem.iter(qn("w:t"), qn("w:lastRenderedPageBreak"),
                          qn("w:footnoteReference")):
        if el.tag == qn("w:t"):
            out.extend(el.text or "")
        elif el.tag == qn("w:lastRenderedPageBreak"):
            out.append("|")
        else:
            out.append("{%d}" % URL_NUM[fn_url[el.get(qn("w:id"))]])
    return out


def _shown(path):
    """Every paragraph of an output file as a string: 'studies {1} {2}.'"""
    doc, _z, ft = load(path)
    fn_url = {f.get(qn("w:id")): fn_text(f) for f in regular(ft)}
    return ["".join(_stream(p._p, fn_url)) for p, _w in _iter(doc)]


_ORACLE_MARK = re.compile(
    r"\[ *\d+(?: *[-\u2013] *\d+)?(?: *[,;] *\d+(?: *[-\u2013] *\d+)?)* *\]")


def _oracle(runs, replace=True, sep=" ", valid=(1, 2, 3, 4), blocked=()):
    """Independent model of the expected output of one paragraph made of
    plain runs (strings; '|' = page-break hint), as a ``_stream`` list.
    ``blocked``: indices k where something that is not plain text (a tab)
    sits between run k-1 and run k.

    Rules: a marker is processed if it lies inside one run, is not a
    reversed range and every number has a URL. Its footnote references stand
    side by side with ``sep`` between them. In replace mode the marker
    disappears, and a processed citation that follows a processed citation
    with only spaces and one comma/semicolon between (and nothing blocking)
    is joined to it: that text disappears too and ``sep`` goes between the
    two groups. A hint inside removed text stays where it was, i.e. just
    before the references that replace it."""
    text, spans, hints = "", [], []
    for spec in runs:
        start = len(text)
        for i, piece in enumerate(spec.split("|")):
            if i:
                hints.append(len(text))
            text += piece
        spans.append((start, len(text)))
    toks = []
    for m in _ORACLE_MARK.finditer(text):
        nums, bad = [], False
        for item in re.split("[,;]", m.group()[1:-1]):
            ends = [int(x) for x in re.findall(r"\d+", item)]
            bad = bad or (len(ends) == 2 and ends[0] > ends[1])
            nums += list(range(ends[0], ends[-1] + 1)) if len(ends) == 2 else ends
        toks.append((m.start(), m.end(), nums, bad,
                     bool(re.fullmatch(r"\[ *\d+ *\]", m.group()))))
    cites, i = [], 0
    while i < len(toks):
        s0, e0, nums, bad, single = toks[i]
        if (single and i + 1 < len(toks) and toks[i + 1][4]
                and re.fullmatch(r"\s*[-\u2013]\s*", text[e0:toks[i + 1][0]])):
            a, b = nums[0], toks[i + 1][2][0]
            cites.append({"start": s0, "end": toks[i + 1][1],
                          "nums": list(range(a, b + 1)), "bad": a > b})
            i += 2
            continue
        cites.append({"start": s0, "end": e0, "nums": nums, "bad": bad})
        i += 1
    prev = None
    for c in cites:
        c["host"] = next((k for k, (a, b) in enumerate(spans)
                          if a <= c["start"] and c["end"] <= b), None)
        c["ok"] = (not c["bad"] and c["host"] is not None
                   and all(n in valid for n in c["nums"]))
        c["joined"] = bool(
            replace and c["ok"] and prev is not None
            and re.fullmatch(r" *[,;]? *", text[prev["end"]:c["start"]])
            and not any(prev["host"] < k <= c["host"] for k in blocked))
        c["from"] = prev["end"] if c["joined"] else c["start"]
        prev = c if c["ok"] else None
    removed = {}                       # offset -> end of the removal it is in
    refs_at = {}
    for c in cites:
        if not c["ok"]:
            continue
        refs = []
        for k, n in enumerate(c["nums"]):
            if sep and (k or c["joined"]):
                refs.extend(sep)
            refs.append("{%d}" % n)
        refs_at[c["end"]] = refs
        if replace:
            for j in range(c["from"], c["end"]):
                removed[j] = c["end"]
    out, deferred = [], {}
    for o in range(len(text) + 1):
        if o in refs_at:
            out.extend(deferred.pop(o, []) + refs_at[o])
        for _ in range(hints.count(o)):
            if o in removed:
                deferred.setdefault(removed[o], []).append("|")
            else:
                out.append("|")
        if o < len(text) and o not in removed:
            out.append(text[o])
    return out


PAGE_BREAK_SPECS = [
    "|head [1] tail",         # the layout Word writes (hint first in run)
    "|[1] tail",              # hint directly before a leading marker
    "head [1]|",              # hint at the very end, run ends with marker
    "he|ad [1] tail",         # hint inside the text before the marker
    "head [|1] tail",         # hint inside the marker itself
    "head [1]| tail",         # hint directly after the marker
    "|a [1]|b| [2] c|",       # several hints and citations in one run
    "x [1]-|[3] y",           # hint inside a range marker
    "x [1,|2] y",             # hint inside a list marker
    "x [1],| [2]| y",         # hints in the text joining two citations
]


@pytest.mark.parametrize("keep", [False, True], ids=["replace", "keep-marker"])
def test_last_rendered_page_break_runs_processed(tmp_path, keep):
    # Word stores <w:lastRenderedPageBreak/> in runs of practically every
    # multi-page document. Such runs used to be rejected as "non-text
    # content" (AMBIGUOUS) - 5 of 30 citations in real test documents.
    doc = Document()
    for spec in PAGE_BREAK_SPECS:
        doc.add_paragraph()._p.append(_run_from_spec(spec))
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = tmp_path / "t.docx"
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path, *(["--keep-marker"] if keep else []))
    assert rc == 0  # includes the validator's structural-count check
    _rows, summary = read_csv(tmp_path)
    assert summary["Ambiguous_citations"] == "0"
    assert summary["Footnotes_inserted"] == "15"   # 6x1 + 2 + 3 (range) + 2 + 2
    shown = _shown(output_path(tmp_path))
    for spec, line in zip(PAGE_BREAK_SPECS, shown):
        assert line == "".join(_oracle([spec], replace=not keep)), spec


def test_citation_in_content_control(tmp_path):
    # Paragraphs of a block-level content control live in w:sdtContent; they
    # used to be invisible (citations neither processed nor reported).
    doc = Document()
    doc.add_paragraph("Intro [1].")
    sdt = OxmlElement("w:sdt")
    sdt.append(OxmlElement("w:sdtPr"))
    content = OxmlElement("w:sdtContent")
    sdt.append(content)
    content.append(doc.add_paragraph("Inside a content control [2].")._p)
    body = doc.element.body
    body.insert(len(body) - 1, sdt)          # before w:sectPr
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = tmp_path / "t.docx"
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc2, _z, ft = load(output_path(tmp_path))
    assert {fn_text(f) for f in regular(ft)} == {"https://example.com/one",
                                                 "https://example.com/two"}
    sdt_text = raw_text_of(doc2.element.body.find(qn("w:sdt")))
    assert sdt_text == "Inside a content control ."
    rows, _summary = read_csv(tmp_path)
    assert any("sdt" in row[1] for row in rows[1:] if row and row[0] == "2")


def test_bold_heading_after_references_ends_list(tmp_path):
    # Documents without heading styles: a bold "Appendix A" after the list
    # used to be glued onto the last reference, so the appendix URL ended
    # up in that reference's footnote.
    doc = Document()
    doc.add_paragraph("Body cites [3].")
    doc.add_paragraph().add_run("References").bold = True
    for r in REFS[:3]:
        doc.add_paragraph(r)
    doc.add_paragraph().add_run("Appendix A").bold = True
    doc.add_paragraph("Survey platform used: https://survey-tool.example/form")
    src = tmp_path / "t.docx"
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _d, _z, ft = load(output_path(tmp_path))
    assert [fn_text(f) for f in regular(ft)] == ["https://example.com/three"]
    _rows, summary = read_csv(tmp_path)
    assert summary["References_detected"] == "3"


def test_table_header_keyword_not_reference_heading(tmp_path):
    # A bold "Reference" column header in an appendix table must not take
    # over as the reference heading (which turned the real reference list
    # into "body" text).
    doc = Document()
    doc.add_paragraph("Body cites [1] and [2].")
    doc.add_paragraph().add_run("References").bold = True
    for r in ["[1] Alpha. https://example.com/one",
              "[2] Beta. https://example.com/two",
              "[3] Gamma. https://example.com/three"]:
        doc.add_paragraph(r)
    doc.add_paragraph().add_run("Appendix B: Studies reviewed").bold = True
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).paragraphs[0].add_run("Reference").bold = True
    tbl.cell(0, 1).paragraphs[0].add_run("Finding").bold = True
    tbl.cell(1, 0).text = "[1]"
    tbl.cell(1, 1).text = "Positive effect"
    src = tmp_path / "t.docx"
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _d, _z, ft = load(output_path(tmp_path))
    assert {fn_text(f) for f in regular(ft)} == {"https://example.com/one",
                                                 "https://example.com/two"}
    _rows, summary = read_csv(tmp_path)
    assert summary["Citation_markers_detected"] == "2"
    assert summary["References_detected"] == "3"


def _field_run(kind=None, instr=None, text=None):
    r = OxmlElement("w:r")
    if kind:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
    elif instr:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = instr
    else:
        el = OxmlElement("w:t")
        el.set(qn("xml:space"), "preserve")
        el.text = text
    r.append(el)
    return r


def test_citation_in_field_result_untouched(tmp_path):
    # EndNote / Zotero / Mendeley Desktop store "[1]" as the RESULT of a
    # complex field. A footnote inserted there is destroyed (or breaks the
    # field) on the next refresh, so it must be left alone and reported.
    doc = Document()
    p = doc.add_paragraph()
    for r in (_field_run(text="Managed by EndNote "),
              _field_run(kind="begin"),
              _field_run(instr=" ADDIN EN.CITE <EndNote><Cite/></EndNote> "),
              _field_run(kind="separate"),
              _field_run(text="[1]"),
              _field_run(kind="end"),
              _field_run(text=" and a typed one [2].")):
        p._p.append(r)
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = tmp_path / "t.docx"
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _d, z, ft = load(output_path(tmp_path))
    assert [fn_text(f) for f in regular(ft)] == ["https://example.com/two"]
    xml = z.read("word/document.xml").decode()
    result = xml[xml.find('w:fldCharType="separate"'):xml.find('w:fldCharType="end"')]
    assert "[1]" in result and "footnoteReference" not in result
    _rows, summary = read_csv(tmp_path)
    assert summary["Ambiguous_citations"] == "1"


def test_uncited_excludes_unprocessed_citations(tmp_path):
    # A reference cited only by a marker that could not be processed
    # (here: split across runs -> AMBIGUOUS) is still cited, not "uncited".
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("See [1] and [3")
    p.add_run("] split.")
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = tmp_path / "t.docx"
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    _rows, summary = read_csv(tmp_path)
    assert summary["Ambiguous_citations"] == "1"
    assert summary["Uncited_references"] == "2, 4"


def test_run_pipeline_returns_actual_output_path(tmp_path):
    import agent_tools
    src = make_doc(tmp_path / "t.docx", ["A [1] B."])
    out_dir = tmp_path / "agent_out"
    code1, out1, rep1 = agent_tools.run_pipeline(str(src), output_dir=str(out_dir))
    code2, out2, rep2 = agent_tools.run_pipeline(str(src), output_dir=str(out_dir))
    assert code1 == 0 and code2 == 0
    assert Path(out1).name == "t_with_footnotes.docx" and Path(out1).is_file()
    assert Path(out2).name == "t_with_footnotes_1.docx" and Path(out2).is_file()
    assert Path(rep2).is_file()
    noref = make_doc(tmp_path / "noref.docx", ["Just prose."], refs=None, heading=None)
    code3, out3, rep3 = agent_tools.run_pipeline(str(noref),
                                                 output_dir=str(tmp_path / "abort"))
    assert code3 == 1 and out3 is None and Path(rep3).is_file()


def test_superscript_inserted_in_schema_order(tmp_path):
    # w:rPr is a strict sequence: w:vertAlign must precede w:lang (present
    # in most Word-authored runs); appending it produced invalid OOXML.
    doc = Document()
    run = doc.add_paragraph().add_run("Tagged run cites [1] here.")
    run.bold = True
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-GB")
    run._r.get_or_add_rPr().append(lang)
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = tmp_path / "t.docx"
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc2, _z, _ft = load(output_path(tmp_path))
    ref = next(doc2.element.body.iter(qn("w:footnoteReference")))
    names = [etree.QName(c).localname for c in ref.getparent().find(qn("w:rPr"))]
    assert names == ["rStyle", "b", "vertAlign", "lang"]


# --------------------------------------------------------------------------
# Several references cited together: [1,2], [1-3], "[1], [2]" -> footnotes
# side by side
# --------------------------------------------------------------------------
def _paragraphs_doc(path, paragraphs):
    """One paragraph per item; an item is a string or a list of run specs
    ('|' = page-break hint), "<tab>" (a tab run), "<bm>" (a bookmark) or
    "<link>" (a hyperlink whose text is "link")."""
    doc = Document()
    for n, item in enumerate(paragraphs):
        p = doc.add_paragraph()
        for spec in ([item] if isinstance(item, str) else item):
            if spec == "<tab>":
                r = OxmlElement("w:r")
                r.append(OxmlElement("w:tab"))
                p._p.append(r)
            elif spec == "<bm>":
                add_bookmark(p, bid=str(100 + n), name=f"bm{n}")
            elif spec == "<link>":
                add_hyperlink(p, "https://example.com/elsewhere", "link")
            else:
                p._p.append(_run_from_spec(spec))
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    doc.save(str(path))
    return Path(path)


GROUP_CASES = [
    # the user's sentences
    ("roundabouts ... published studies [1,2].", "roundabouts ... published studies {1} {2}."),
    ("behavior of drivers [3,4].", "behavior of drivers {3} {4}."),
    # several numbers in one bracket
    ("A [1, 2] B", "A {1} {2} B"),
    ("A [1;3] B", "A {1} {3} B"),
    ("A [ 2 , 4 ] B", "A {2} {4} B"),
    ("A [1-3] B", "A {1} {2} {3} B"),
    ("A [1\u20133] B", "A {1} {2} {3} B"),
    ("A [1, 3-4] B", "A {1} {3} {4} B"),
    ("A [1]-[3] B", "A {1} {2} {3} B"),
    # separate brackets that belong together
    ("A [1], [2] B", "A {1} {2} B"),
    ("A [1],[2] B", "A {1} {2} B"),
    ("A [1][2] B", "A {1} {2} B"),
    ("A [1] [2] B", "A {1} {2} B"),
    ("A [1]; [2] B", "A {1} {2} B"),
    ("A [1,2], [3] B", "A {1} {2} {3} B"),
    ("A [1]-[3], [4] B", "A {1} {2} {3} {4} B"),
    # words and other text are never removed
    ("A [1] and [2] B", "A {1} and {2} B"),
    ("A [1], [2], and [3] B", "A {1} {2}, and {3} B"),
    ("A [1],, [2] B", "A {1},, {2} B"),
    # problems: nothing partial, nothing guessed
    ("A [1,9] B", "A [1,9] B"),                 # 9 is not in the list
    ("A [1], [9] B", "A {1}, [9] B"),           # [1] on its own, [9] kept
    ("A [9], [1] B", "A [9], {1} B"),
    ("A [3-1] B", "A [3-1] B"),                 # reversed range
    ("A [1-400] B", "A [1-400] B"),             # not a citation
    ("A [1, p. 4] B", "A [1, p. 4] B"),         # not a citation either
]


def test_citations_with_several_references(tmp_path):
    src = _paragraphs_doc(tmp_path / "t.docx", [given for given, _ in GROUP_CASES])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0                                # validation passed
    shown = _shown(output_path(tmp_path))
    for (given, wanted), line in zip(GROUP_CASES, shown):
        assert line == wanted, given
    rows, summary = read_csv(tmp_path)
    assert summary["Unresolved_citations"] == "3"   # [1,9] and both [9]
    statuses = {(r[2], r[6]) for r in rows[1:] if len(r) > 6}
    assert ("[1,9]", "ERROR") in statuses and ("[3-1]", "AMBIGUOUS") in statuses
    assert ("[1-400]", "AMBIGUOUS") in statuses


def test_keep_marker_mode_with_several_references(tmp_path):
    src = _paragraphs_doc(tmp_path / "t.docx", [
        "studies [1,2].", "A [1], [2] B", "A [1-3] B"])
    rc, _ = run_tool(src, tmp_path, "--keep-marker")
    assert rc == 0
    assert _shown(output_path(tmp_path))[:3] == [
        "studies [1,2]{1} {2}.", "A [1]{1}, [2]{2} B", "A [1-3]{1} {2} {3} B"]
    # a second run finds the existing footnotes and adds nothing
    rc2, out2 = run_tool(output_path(tmp_path), tmp_path / "r2", "--keep-marker")
    assert rc2 == 0
    _rows, summary = read_csv(tmp_path / "r2")
    assert summary["Footnotes_inserted"] == "0"
    assert summary["Previously_processed"] == "4"


@pytest.mark.parametrize("sep, wanted", [
    (",", ["studies {1},{2}.", "A {1},{2} B"]),
    ("", ["studies {1}{2}.", "A {1}{2} B"]),
])
def test_separator_option(tmp_path, sep, wanted):
    src = _paragraphs_doc(tmp_path / "t.docx", ["studies [1,2].", "A [1], [2] B"])
    rc, _ = run_tool(src, tmp_path, "--separator", sep)
    assert rc == 0
    assert _shown(output_path(tmp_path))[:2] == wanted


def test_separator_is_formatted_like_the_footnote_references(tmp_path):
    src = make_doc(tmp_path / "t.docx", [["Plain ", ("bold claim [1,2]", True), "."]])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    doc, _z, _ft = load(output_path(tmp_path))
    ref = next(doc.element.body.iter(qn("w:footnoteReference"))).getparent()
    sep = ref.getnext()
    assert "".join(t.text for t in sep.iter(qn("w:t"))) == " "
    names = [etree.QName(c).localname for c in sep.find(qn("w:rPr"))]
    assert names == [etree.QName(c).localname for c in ref.find(qn("w:rPr"))]
    assert names == ["rStyle", "b", "vertAlign"]   # bold kept, superscript added
    assert sep.getnext().find(qn("w:footnoteReference")) is not None


def test_joining_across_invisible_run_splits(tmp_path):
    # Word splits text into runs for editing history, spell checking, ...
    # "[1], [2]" can be several runs; the ", " between may be its own run.
    src = _paragraphs_doc(tmp_path / "t.docx", [
        ["A [1]", ", ", "[2] B"],
        ["A [1],", " [2] B"],
        ["A [1]", ",|", " [2] B"],        # page-break hint in the removed text
        ["A [1]", "<bm>", ", [2] B"],     # a bookmark between is fine
        ["A [1]", "<tab>", ", [2] B"],    # a tab is real content: not joined
        ["A [1]", "<link>", ", [2] B"],   # so is a link (its text is not in the
                                          # run text the markers are found in)
        ["A [1", "], [2] B"],             # [1] split by formatting: left alone
    ])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0                         # incl. structural checks (bookmark, hint)
    assert _shown(output_path(tmp_path))[:7] == [
        "A {1} {2} B", "A {1} {2} B", "A {1}| {2} B", "A {1} {2} B",
        "A {1}, {2} B", "A {1}link, {2} B", "A [1], {2} B"]


def test_never_joined_across_an_existing_footnote(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("A [1]")
    add_preexisting_footnote(doc, p, "My own note.")
    p.add_run(", [2] B")
    doc.add_heading("References", level=1)
    for r in REFS:
        doc.add_paragraph(r)
    src = tmp_path / "t.docx"
    doc.save(str(src))
    rc, _ = run_tool(src, tmp_path)
    assert rc == 0
    out_doc, _z, ft = load(output_path(tmp_path))
    assert raw_text_of(next(iter_paragraphs(out_doc))._p) == "A ,  B"   # ", " kept
    assert len(body_ref_ids(out_doc)) == 3 and len(regular(ft)) == 3


def test_validator_catches_text_removed_between_citations(tmp_path, monkeypatch):
    # The validator checks the output on its own: if joining ever removed
    # more than spaces and a comma, validation must fail.
    import re as _re
    import word_footnotes
    monkeypatch.setattr(word_footnotes, "JOIN_GAP_RE", _re.compile(r".*"))
    src = _paragraphs_doc(tmp_path / "t.docx", ["A [1] and [2] B"])
    rc, _ = run_tool(src, tmp_path)
    assert rc == 1
    assert _shown(output_path(tmp_path))[0] == "A {1} {2} B"   # "and" was lost
    _rows, summary = read_csv(tmp_path)
    assert summary["Status"] == "VALIDATION FAILED"


def test_real_world_list_citations_are_reported(tmp_path):
    rows_src = _paragraphs_doc(tmp_path / "t.docx", ["audit quality [3,4], or growth [2]."])
    rc, _ = run_tool(rows_src, tmp_path)
    assert rc == 0
    rows, summary = read_csv(tmp_path)
    assert summary["Citation_markers_detected"] == "2"
    assert summary["Footnotes_inserted"] == "3"
    listed = [r for r in rows[1:] if len(r) > 7 and r[2] == "[3,4]"]
    assert [r[0] for r in listed] == ["3", "4"]
    assert all(r[6] == "SUCCESS" and "Several references" in r[7] for r in listed)


def _random_case(rng):
    """A random paragraph full of citations, hints and run splits."""
    def num():
        return rng.choice([1, 2, 3, 4, 4, 9])
    forms = [
        lambda: f"[{num()}]",
        lambda: f"[{num()},{num()}]",
        lambda: f"[{num()}, {num()}]",
        lambda: f"[{num()};{num()}]",
        lambda: f"[{num()}-{num()}]",
        lambda: f"[{num()} - {num()}]",
        lambda: f"[{num()}]-[{num()}]",
        lambda: f"[{num()}, {num()}-{num()}]",
    ]
    text = ""
    for _ in range(rng.randint(1, 4)):
        text += rng.choice(["Text ", "a b ", "word", "x, ", "y; ", "", " "])
        group = [rng.choice(forms)() for _ in range(rng.randint(1, 3))]
        gaps = [rng.choice([", ", ",", " ", "", "; ", " and ", ",, ", " ,  "])
                for _ in group[1:]]
        text += group[0] + "".join(g + m for g, m in zip(gaps, group[1:]))
    text += rng.choice([".", " end.", ""])
    for _ in range(rng.randint(0, 2)):                    # page-break hints
        i = rng.randint(0, len(text))
        text = text[:i] + "|" + text[i:]
    cuts = sorted(rng.sample(range(len(text) + 1), rng.randint(0, 3)))
    runs = [text[a:b] for a, b in zip([0] + cuts, cuts + [len(text)])]
    between = [rng.choice([None, None, None, "<tab>", "<bm>"]) for _ in runs[1:]]
    items, blocked = [runs[0]], set()
    for k, (extra, spec) in enumerate(zip(between, runs[1:]), 1):
        if extra:
            items.append(extra)
        if extra == "<tab>":
            blocked.add(k)
        items.append(spec)
    return items, runs, blocked


@pytest.mark.parametrize("keep, sep", [(False, " "), (False, ","), (True, " "), (False, "")],
                         ids=["replace", "comma", "keep-marker", "no-separator"])
def test_random_paragraphs_match_an_independent_model(tmp_path, keep, sep):
    import random
    rng = random.Random(20261002 + len(sep) + keep)
    cases = [_random_case(rng) for _ in range(150)]
    src = _paragraphs_doc(tmp_path / "t.docx", [items for items, _r, _b in cases])
    args = ["--separator", sep] + (["--keep-marker"] if keep else [])
    rc, _ = run_tool(src, tmp_path, *args)
    assert rc == 0                                         # validator agrees
    shown = _shown(output_path(tmp_path))
    for (items, runs, blocked), line in zip(cases, shown):
        assert line == "".join(_oracle(runs, replace=not keep, sep=sep,
                                       blocked=blocked)), items


# --------------------------------------------------------------------------
# The Windows COM backend, driven through a small fake of Word's object model
# (Word itself only exists on Windows; this checks the backend's logic)
# --------------------------------------------------------------------------
class _FakeFont:
    def __init__(self, rng):
        self._rng = rng

    @property
    def Superscript(self):
        return all(c[1] for c in self._rng.doc.chars[self._rng.Start:self._rng.End])

    @Superscript.setter
    def Superscript(self, value):
        for c in self._rng.doc.chars[self._rng.Start:self._rng.End]:
            c[1] = value


class _FakeFind:
    def __init__(self, rng):
        self._rng = rng
        self.Text = ""

    def ClearFormatting(self):
        pass

    def Execute(self):
        i = self._rng.doc.text().find(self.Text, self._rng.Start, self._rng.End)
        if i < 0:
            return False
        self._rng.Start, self._rng.End = i, i + len(self.Text)
        return True


class _FakeRange:
    def __init__(self, doc, start, end):
        self.doc, self.Start, self.End = doc, start, end
        self.Find = _FakeFind(self)
        self.Font = _FakeFont(self)

    @property
    def Text(self):
        return self.doc.text()[self.Start:self.End]

    def Delete(self):
        del self.doc.chars[self.Start:self.End]
        self.End = self.Start

    def InsertAfter(self, text):
        self.doc.chars[self.End:self.End] = [[ch, False] for ch in text]
        self.End += len(text)


class _FakeDoc:
    """Characters are [char, superscript]; a footnote reference is the
    character ("FN", url) - shown as '\\x02' in Text, as Word does."""

    def __init__(self, paragraphs):
        self.chars = [[ch, False] for ch in "\r".join(paragraphs) + "\r"]
        self.saved_as = None
        self.Footnotes = self

    def text(self):
        return "".join("\x02" if isinstance(c[0], tuple) else c[0] for c in self.chars)

    def Range(self, start, end):
        return _FakeRange(self, start, end)

    @property
    def Content(self):
        return _FakeRange(self, 0, len(self.chars))

    def Add(self, Range, Text):                      # doc.Footnotes.Add
        self.chars.insert(Range.Start, [("FN", Text), True])
        import types
        return types.SimpleNamespace(Reference=_FakeRange(self, Range.Start, Range.Start + 1))

    def SaveAs(self, path):
        self.saved_as = path

    def Close(self, SaveChanges=False):
        pass

    def shown(self):
        """'studies {1}^ ^{2}.' - '^' marks superscript characters."""
        out = []
        for ch, sup in self.chars:
            if isinstance(ch, tuple):
                out.append("{%d}" % URL_NUM[ch[1]])
            else:
                out.append(ch + ("^" if sup else ""))
        return "".join(out).split("\r")


def test_com_backend_logic_with_a_fake_word(tmp_path, monkeypatch):
    import types
    import word_footnotes
    from citation_detector import CitationDetector
    from document_reader import DocumentReader
    from config import Config
    paragraphs = ["studies [1,2].", "A [1], [2] B", "A [1] and [3-4] B",
                  "A [1], [9] B", "A [2], [1] [3]"]
    src = _paragraphs_doc(tmp_path / "t.docx", paragraphs)
    paras = DocumentReader(src).all_paragraphs()
    start = ReferenceParser(Config()).find_reference_section(paras, Config())[0]
    ref_map, _w = ReferenceParser().parse_references(paras, start)
    cites = CitationDetector().detect(paras, start)
    fake = _FakeDoc([raw_text_of(p._p) for p in paras])
    word = types.SimpleNamespace(Visible=True, Quit=lambda: None,
                                 Documents=types.SimpleNamespace(Open=lambda path: fake))
    client = types.ModuleType("win32com.client")
    client.Dispatch = lambda name: word
    pkg = types.ModuleType("win32com")
    pkg.client = client
    monkeypatch.setitem(sys.modules, "win32com", pkg)
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setattr(word_footnotes, "os", types.SimpleNamespace(
        name="nt", path=word_footnotes.os.path))
    word_footnotes.process_document_com(str(src), cites, ref_map, True,
                                        str(tmp_path / "out.docx"), separator=" ")
    assert fake.saved_as == str((tmp_path / "out.docx").resolve())
    assert fake.shown()[:5] == [
        "studies {1} ^{2}.",            # the separator is superscript like the refs
        "A {1} ^{2} B",                 # ", " removed: joined
        "A {1} and {3} ^{4} B",         # words stay
        "A {1}, [9] B",                 # [9] has no reference: kept, not joined
        "A {2} ^{1} ^{3}",
    ]
    by_text = {(c.para_index, c.marker_text): c for c in cites}
    assert by_text[(1, "[2]")].joined and by_text[(1, "[2]")].join_from == 5
    assert by_text[(3, "[9]")].status == "ERROR" and not by_text[(3, "[1]")].joined
