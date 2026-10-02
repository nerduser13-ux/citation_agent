"""tests/test_agent.py

Tests for the AI-agent layer: link_checker.py, agent_tools.Toolbox and
agent.py. Everything runs offline:

* registries (Crossref, Europe PMC, doi.org) and web pages are served by a
  local fake server;
* "Gemini" is a local server speaking the real Gemini REST format, driven
  through Google's real google-genai SDK (skipped if it isn't installed).

Run:  python -m pytest tests/ -v
"""
import hashlib
import json
import socket
import sys
import threading
import urllib.parse
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import link_checker as lc                                   # noqa: E402
from agent_tools import Toolbox                             # noqa: E402

try:
    import google.genai  # noqa: F401
    HAVE_GENAI = sys.version_info >= (3, 10)
except ImportError:
    HAVE_GENAI = False
needs_genai = pytest.mark.skipif(
    not HAVE_GENAI, reason="google-genai not installed (pip install -r requirements-agent.txt)")


# --------------------------------------------------------------------------
# Fake HTTP server (registries, web pages, Gemini)
# --------------------------------------------------------------------------
class FakeServer:
    """GET: ``routes`` {path: (status, content_type, body[, headers]) or
    callable(query, headers) -> tuple}. POST: answered from ``post_script`` in
    order (item = response dict, (status, dict[, headers]), or
    callable(headers) -> one of those)."""

    def __init__(self, routes=None, post_script=None):
        self.routes = routes or {}
        self.post_script = list(post_script or [])
        self.gets, self.posts = [], []
        server = self

        class Handler(BaseHTTPRequestHandler):
            def _send(self, status, ctype, body, headers=None):
                if isinstance(body, (dict, list)):
                    body = json.dumps(body)
                if isinstance(body, str):
                    body = body.encode("utf-8")
                self.send_response(status)
                self.send_header("Content-Type", ctype)
                self.send_header("Content-Length", str(len(body)))
                for k, v in (headers or {}).items():
                    self.send_header(k, v)
                self.end_headers()
                self.wfile.write(body)

            def do_GET(self):
                parsed = urllib.parse.urlparse(self.path)
                server.gets.append(self.path)
                route = server.routes.get(urllib.parse.unquote(parsed.path))
                if callable(route):
                    route = route(urllib.parse.parse_qs(parsed.query), dict(self.headers))
                if route is None:
                    return self._send(404, "text/plain", "not found")
                self._send(*route)

            def do_POST(self):
                length = int(self.headers.get("Content-Length") or 0)
                body = json.loads(self.rfile.read(length) or b"{}")
                server.posts.append({"path": self.path, "headers": dict(self.headers),
                                     "body": body})
                if not server.post_script:
                    return self._send(500, "application/json", {"error": {
                        "code": 500, "message": "script exhausted", "status": "INTERNAL"}})
                item = server.post_script.pop(0)
                if callable(item):
                    item = item(dict(self.headers))
                status, payload, *extra = item if isinstance(item, tuple) else (200, item)
                self._send(status, "application/json", payload, *extra)

            def log_message(self, *args):
                pass

        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.base = f"http://127.0.0.1:{self.httpd.server_port}"
        threading.Thread(target=self.httpd.serve_forever, kwargs={"poll_interval": 0.05},
                         daemon=True).start()

    def close(self):
        self.httpd.shutdown()
        self.httpd.server_close()


def _closed_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


JSON = "application/json"
HTML = "text/html; charset=utf-8"
AI_TITLE = "Artificial intelligence adoption in SMEs"
REF_AI = ("3. S\u00e1nchez E, Calder\u00f3n R. Artificial intelligence adoption in SMEs: "
          "survey based on TOE\u2013DOI framework. Appl Sci. 2025;15(12):6465.")


def crossref(title, subtitle=None):
    rec = {"title": [title], "issued": {"date-parts": [[2025, 6]]}}
    if subtitle:
        rec["subtitle"] = [subtitle]
    return (200, JSON, {"status": "ok", "message": rec})


def _crossref_filter(query, headers=None):
    if query.get("filter") == ["alternative-id:S2444569X25000320"]:
        return (200, JSON, {"message": {"items": [
            {"title": [AI_TITLE], "DOI": "10.1016/j.jik.2025.100687"}]}})
    return (200, JSON, {"message": {"items": []}})


def _europepmc(query, headers=None):
    if query.get("query") == ["PMCID:PMC12749562"]:
        return (200, JSON, {"hitCount": 1, "resultList": {"result": [
            {"title": AI_TITLE + ".", "pubYear": "2025", "doi": "10.1/x"}]}})
    return (200, JSON, {"hitCount": 0, "resultList": {"result": []}})


PDF_XMP = (b"%PDF-1.6\n<x:xmpmeta><rdf:RDF><rdf:Description><dc:title><rdf:Alt>"
           b"<rdf:li xml:lang=\"x-default\">Artificial intelligence adoption in SMEs"
           b"</rdf:li></rdf:Alt></dc:title></rdf:Description></rdf:RDF></x:xmpmeta>\n")

ROUTES = {
    "/crossref/works/10.1234/good": crossref(
        "Artificial Intelligence Adoption in SMEs", "Survey Based on TOE-DOI Framework"),
    "/crossref/works/10.1234/other": crossref("Deep learning for protein structure prediction"),
    "/crossref/works/10.1234/partial": crossref(
        "Artificial intelligence and blockchain adoption in large manufacturing firms"),
    "/crossref/works/10.1108/IJSE-12-2021-0752": crossref(AI_TITLE),
    "/crossref/works": _crossref_filter,
    "/handles/10.9999/nothing": (404, JSON, {"responseCode": 100}),
    "/europepmc/search": _europepmc,
    "/page/good": (200, HTML, "<html><head><title>Journal site</title>"
                   f"<meta name=\"citation_title\" content=\"{AI_TITLE}\"></head>"
                   "<body><title>not this</title></body></html>"),
    "/page/sitename": (200, HTML, f"<html><head><title>{AI_TITLE} | Journal of "
                       "Things</title></head></html>"),
    "/page/unrelated": (200, HTML, "<title>Welcome to our homepage</title>"),
    "/page/soft404": (200, HTML, "<title>Page not found</title>"),
    "/page/blocked": (403, HTML, "<title>Forbidden</title>"),
    "/page/robot": (200, HTML, "<title>Just a moment...</title>"),
    "/page/pdf-xmp": (200, "application/pdf", PDF_XMP),
    "/page/pdf-plain": (200, "application/pdf", b"%PDF-1.4\n1 0 obj<<>>endobj\n"),
    "/page/moved": (301, HTML, "", {"Location": "/page/good"}),
}


@pytest.fixture
def web(monkeypatch):
    srv = FakeServer(dict(ROUTES))
    monkeypatch.setattr(lc, "CROSSREF_API", srv.base + "/crossref")
    monkeypatch.setattr(lc, "EUROPEPMC_API", srv.base + "/europepmc")
    monkeypatch.setattr(lc, "DOI_HANDLE_API", srv.base + "/handles")
    monkeypatch.setattr(lc, "TIMEOUT", 5)
    yield srv
    srv.close()


# --------------------------------------------------------------------------
# link_checker: pure functions
# --------------------------------------------------------------------------
@pytest.mark.parametrize("url, expected", [
    ("https://doi.org/10.3390/app15126465", "10.3390/app15126465"),
    ("http://dx.doi.org/10.1080/09537287.2022.2131620", "10.1080/09537287.2022.2131620"),
    ("https://link.springer.com/article/10.1007/s11156-024-01278-0", "10.1007/s11156-024-01278-0"),
    ("https://www.tandfonline.com/doi/full/10.1080/09537287.2022.2131620",
     "10.1080/09537287.2022.2131620"),
    ("https://link.springer.com/content/pdf/10.1007/s11156-024-01278-0.pdf",
     "10.1007/s11156-024-01278-0"),
    ("https://example.org/search?doi=10.5267%2Fj.ijdns.2023.12.021", "10.5267/j.ijdns.2023.12.021"),
    ("https://www.mdpi.com/2076-3417/15/12/6465", None),
    ("https://pmc.ncbi.nlm.nih.gov/articles/PMC12749562/", None),
    ("https://www.growingscience.com/ijds/Vol7/ijdns_2022_120.pdf", None),
])
def test_doi_extraction(url, expected):
    cands = lc.doi_candidates(url)
    assert (cands[0] if cands else None) == expected


def test_doi_trailing_publisher_path_is_trimmed():
    cands = lc.doi_candidates(
        "https://www.emerald.com/insight/content/doi/10.1108/IJSE-12-2021-0752/full/html")
    assert cands[0] == "10.1108/IJSE-12-2021-0752/full/html"
    assert "10.1108/IJSE-12-2021-0752" in cands
    # doi.org links are taken literally - a DOI may contain slashes
    assert lc.doi_candidates("https://doi.org/10.1002/a/b") == ["10.1002/a/b"]


def test_pii_and_pmcid_extraction():
    assert lc.extract_pii(
        "https://www.sciencedirect.com/science/article/pii/S2444569X25000320") == "S2444569X25000320"
    assert lc.extract_pii("https://www.cell.com/fulltext/S0092-8674(20)30230-2") is None
    assert lc.extract_pmcid("https://pmc.ncbi.nlm.nih.gov/articles/PMC12749562/") == "PMC12749562"
    assert lc.extract_pmcid("https://example.com/paper") is None


def test_html_title_priority_and_entities():
    page = (b"<html><head><title>Site | Home</title>"
            b"<meta property=\"og:title\" content=\"OG title\">"
            b"<meta name=\"citation_title\" content=\"Caf&eacute; &amp; <i>AI</i>\">"
            b"</head><body><meta name=\"citation_title\" content=\"late\"></body></html>")
    assert lc.html_title(page, HTML) == "Caf\u00e9 & AI"
    assert lc.html_title(b"<title>  Only\n title </title>") == "Only title"
    assert lc.html_title(b"<p>no title</p>") == ""


def test_pdf_title():
    assert lc.pdf_title(PDF_XMP) == AI_TITLE
    assert lc.pdf_title(b"%PDF-1.4 << /Title (Paper \\(draft\\)) >>") == "Paper (draft)"
    utf16 = "\u00c9tude".encode("utf-16-be")
    assert lc.pdf_title(b"%PDF << /Title (\xfe\xff" + utf16 + b") >>") == "\u00c9tude"
    assert lc.pdf_title(b"%PDF-1.4 nothing") == ""


def test_title_match():
    assert lc.title_match("Artificial Intelligence Adoption in SMEs", REF_AI) == 1.0
    # accents, dashes and case are normalised
    assert lc.title_match("Survey based on TOE-DOI framework", REF_AI) == 1.0
    assert lc.title_match("Deep learning for protein structure prediction", REF_AI) == 0.0
    assert lc.title_match("", REF_AI) == 0.0
    assert lc._page_title_match(f"{AI_TITLE} | Journal of Things", REF_AI) == 1.0


# --------------------------------------------------------------------------
# link_checker: against the fake registries / websites
# --------------------------------------------------------------------------
@pytest.mark.parametrize("url, verdict, source", [
    ("https://doi.org/10.1234/good", lc.OK, "Crossref"),
    ("https://doi.org/10.1234/other", lc.PROBLEM, "Crossref"),
    ("https://doi.org/10.1234/partial", lc.CHECK, "Crossref"),
    ("https://www.emerald.com/insight/content/doi/10.1108/IJSE-12-2021-0752/full/html",
     lc.OK, "Crossref"),
    ("https://www.sciencedirect.com/science/article/pii/S2444569X25000320", lc.OK, "Crossref"),
    ("https://pmc.ncbi.nlm.nih.gov/articles/PMC12749562/", lc.OK, "Europe PMC"),
])
def test_registry_checks(web, url, verdict, source):
    res = lc.check_link(url, REF_AI)
    assert res["verdict"] == verdict, res
    assert res["source"] == source
    assert res["found_title"]


def test_registry_result_details(web):
    res = lc.check_link("https://doi.org/10.1234/good", REF_AI)
    assert res["found_title"] == ("Artificial Intelligence Adoption in SMEs: "
                                  "Survey Based on TOE-DOI Framework")
    assert res["match"] == 1.0 and res["found_year"] == 2025 and res["doi"] == "10.1234/good"
    wrong = lc.check_link("https://doi.org/10.1234/other", REF_AI)
    assert "different paper" in wrong["detail"]
    assert "Deep learning for protein structure prediction" in wrong["detail"]


def test_nonexistent_doi_is_a_problem(web):
    res = lc.check_link("https://doi.org/10.9999/nothing", REF_AI)
    assert res["verdict"] == lc.PROBLEM
    assert "does not exist" in res["detail"]


@pytest.mark.parametrize("page, verdict", [
    ("good", lc.OK),           # citation_title meta wins over <title>
    ("sitename", lc.OK),       # "Title | Journal" - best segment
    ("moved", lc.OK),          # redirects are followed
    ("pdf-xmp", lc.OK),
    ("pdf-plain", lc.CHECK),   # opens, title unreadable
    ("unrelated", lc.CHECK),   # website titles never prove a "different paper"
    ("soft404", lc.PROBLEM),
    ("missing", lc.PROBLEM),   # HTTP 404
    ("blocked", lc.UNVERIFIED),
    ("robot", lc.UNVERIFIED),
])
def test_page_checks(web, page, verdict):
    res = lc.check_link(f"{web.base}/page/{page}", REF_AI)
    assert res["verdict"] == verdict, res


def test_check_references_batch(web):
    refs = {
        1: {"text": REF_AI, "url": "https://doi.org/10.1234/good"},
        2: {"text": REF_AI, "url": f"{web.base}/page/blocked"},
        3: {"text": "3. No link here. 2020.", "url": None},
        4: {"text": REF_AI, "url": "http://no-such-host.invalid/paper"},
    }
    report = lc.check_references(refs)
    verdicts = {r["reference"]: r["verdict"] for r in report["results"]}
    # the .invalid address can't resolve while other links worked -> gone
    assert verdicts == {1: lc.OK, 2: lc.UNVERIFIED, 3: lc.NO_LINK, 4: lc.PROBLEM}
    assert report["summary"] == {"OK": 1, "CHECK": 0, "PROBLEM": 1,
                                 "UNVERIFIED": 1, "NO_LINK": 1}
    assert report["checked"] == 4 and report["note"] == ""
    assert all(set(r) >= {"reference", "reference_text", "url", "verdict", "detail"}
               for r in report["results"])
    only = lc.check_references(refs, numbers=[3])
    assert [r["reference"] for r in only["results"]] == [3]


def test_offline_is_unverified_not_broken(monkeypatch):
    dead = f"http://127.0.0.1:{_closed_port()}"
    for attr in ("CROSSREF_API", "EUROPEPMC_API", "DOI_HANDLE_API"):
        monkeypatch.setattr(lc, attr, dead)
    refs = {1: {"text": REF_AI, "url": dead + "/paper"},
            2: {"text": REF_AI, "url": "http://no-such-host.invalid/paper"}}
    report = lc.check_references(refs)
    assert [r["verdict"] for r in report["results"]] == [lc.UNVERIFIED, lc.UNVERIFIED]
    assert "internet" in report["note"]


# --------------------------------------------------------------------------
# Toolbox
# --------------------------------------------------------------------------
def make_doc(path, body, refs, heading="References"):
    doc = Document()
    for text in body:
        doc.add_paragraph(text)
    if heading:
        doc.add_heading(heading, level=1)
    for ref in refs:
        doc.add_paragraph(ref)
    doc.save(str(path))
    return path


REFS = ["1. Alpha A. First paper. 2021. https://example.com/one",
        "2. Beta B. Second paper. 2022. https://example.com/two",
        "3. Gamma C. Third paper. 2023. https://example.com/three"]


@pytest.fixture
def box(tmp_path):
    inp = tmp_path / "input"
    inp.mkdir()
    make_doc(inp / "My Essay.docx", ["Intro [1] and [2].", "More [3]."], REFS)
    make_doc(inp / "Report.docx", ["Only [1]."], REFS)
    (inp / "~$My Essay.docx").write_bytes(b"lock")       # Word's lock file
    (inp / "notes.txt").write_text("x")
    return Toolbox(inp, tmp_path / "output")


def _sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def test_list_documents_skips_lock_and_other_files(box):
    res = box.call("list_documents")
    assert res["documents"] == ["My Essay.docx", "Report.docx"]
    assert res["output_files"] == []


@pytest.mark.parametrize("name, expected", [
    ("My Essay.docx", "My Essay.docx"),
    ("my essay", "My Essay.docx"),
    ("essay", "My Essay.docx"),
    ("REPORT.DOCX", "Report.docx"),
    ("C:\\Users\\me\\citation_agent\\input\\Report.docx", "Report.docx"),
])
def test_resolve_names(box, name, expected):
    assert box.resolve(name).name == expected


@pytest.mark.parametrize("name, fragment", [
    ("../../etc/passwd", "no document called"),
    ("..\\main.py", "no document called"),
    ("~$My Essay.docx", "no document called"),
    ("e", "matches several"),
    ("", "no document name"),
])
def test_resolve_refuses_everything_else(box, name, fragment):
    res = box.call("analyze_document", {"name": name})
    assert fragment in res["error"].lower()


def test_analyze_is_a_preview_that_saves_nothing(box):
    before = _sha(box.input_dir / "My Essay.docx")
    res = box.call("analyze_document", {"name": "essay"})
    assert res["document"] == "My Essay.docx" and res["preview_only"] is True
    assert res["citations_found"] == 3
    assert res["footnotes_that_would_be_added"] == 3
    assert res["problems"] == []
    assert [r["number"] for r in res["references"]] == [1, 2, 3]
    assert res["references"][0]["url"] == "https://example.com/one"
    assert not box.output_dir.exists()                  # nothing written
    assert _sha(box.input_dir / "My Essay.docx") == before


def test_add_footnotes_creates_copy_and_keeps_original(box):
    src = box.input_dir / "My Essay.docx"
    before = _sha(src)
    res = box.call("add_footnotes", {"name": "essay"})
    assert res["done"] is True and res["validation"] == "PASS"
    assert res["footnotes_added"] == 3 and res["original_unchanged"] is True
    assert res["output_file"] == "My Essay_with_footnotes.docx"
    assert (box.output_dir / res["output_file"]).is_file()
    assert _sha(src) == before
    again = box.call("add_footnotes", {"name": "essay", "keep_marker": True})
    assert again["output_file"] == "My Essay_with_footnotes_1.docx"   # never overwrites
    assert box.call("list_documents")["output_files"] == [
        "My Essay_with_footnotes.docx", "My Essay_with_footnotes_1.docx",
        "citation_review_report.csv"]


def test_add_footnotes_puts_grouped_citations_side_by_side(box):
    from docx import Document as _Document
    make_doc(box.input_dir / "Group.docx", ["Studies [1,2] and [1], [3]."], REFS)
    res = box.call("add_footnotes", {"name": "Group"})
    assert res["footnotes_added"] == 4 and res["validation"] == "PASS"
    assert res["problems"] == []
    out = _Document(str(box.output_dir / res["output_file"]))
    assert out.paragraphs[0].text == "Studies   and  ."     # refs + spaces between
    res = box.call("add_footnotes", {"name": "Group", "separator": "comma"})
    out = _Document(str(box.output_dir / res["output_file"]))
    assert out.paragraphs[0].text == "Studies , and ,." and res["validation"] == "PASS"
    assert "separator must be" in box.call(
        "add_footnotes", {"name": "Group", "separator": ";"})["error"]
    spec = next(t for t in box.TOOL_SPECS if t["name"] == "add_footnotes")
    assert spec["parameters"]["properties"]["separator"]["enum"] == ["space", "comma", "none"]


def test_problems_are_explained(box):
    make_doc(box.input_dir / "Broken.docx", ["A [1]. B [5]. C [2]."],
             ["1. Alpha A. First paper. 2021. https://example.com/one",
              "2. Beta B. Second paper without a link. 2022.",
              "3. Gamma C. Third paper. 2023. https://example.com/three"])
    res = box.call("add_footnotes", {"name": "Broken"})
    assert res["footnotes_added"] == 1 and res["validation"] == "PASS"
    assert res["references_without_link"] == [2]
    by_marker = {p["citation"]: p for p in res["problems"]}
    assert set(by_marker) == {"[5]", "[2]"}
    assert all(p["status"] == "ERROR" and p["meaning"] for p in res["problems"])
    assert "not found" in by_marker["[5]"]["why"].lower()
    assert "section" in by_marker["[5]"]["where"]


def test_no_reference_list_is_explained(box):
    make_doc(box.input_dir / "NoRefs.docx", ["Text [1]."], [], heading=None)
    res = box.call("add_footnotes", {"name": "NoRefs"})
    assert res["reference_list_found"] is False and "References" in res["message"]
    assert not list(box.output_dir.glob("NoRefs*.docx"))
    assert box.call("check_links", {"name": "NoRefs"})["reference_list_found"] is False


def test_check_links_tool(box, web):
    make_doc(box.input_dir / "Links.docx", ["Text [1], [2] and [3]."], [
        f"1. S\u00e1nchez E. {AI_TITLE}. 2025. {web.base}/page/good",
        "2. Other A. Some unrelated title. 2024. https://doi.org/10.1234/other",
        "3. No Link B. A reference without any link. 2023."])
    res = box.call("check_links", {"name": "Links"})
    assert res["document"] == "Links.docx"
    assert [(r["reference"], r["verdict"]) for r in res["results"]] == [
        (1, "OK"), (2, "PROBLEM"), (3, "NO_LINK")]
    # JSON numbers may arrive as floats; unknown numbers are reported
    some = box.call("check_links", {"name": "Links", "numbers": [2.0, 9]})
    assert [r["reference"] for r in some["results"]] == [2]
    assert some["not_in_reference_list"] == [9]


def test_call_rejects_unknown_tools_and_bad_arguments(box):
    assert "Unknown tool" in box.call("delete_everything", {})["error"]
    assert "error" in box.call("check_links", {"name": "essay", "numbers": ["x"]})
    # unexpected arguments are ignored, not passed through
    assert box.call("list_documents", {"path": "/etc"})["documents"]
    assert "error" in box.call("analyze_document", {})


def test_tool_results_are_json_serialisable(box, web):
    for name, args in [("list_documents", {}), ("analyze_document", {"name": "essay"}),
                       ("add_footnotes", {"name": "Report"})]:
        json.dumps(box.call(name, args))


# --------------------------------------------------------------------------
# agent.py with a fake Gemini
# --------------------------------------------------------------------------
def fc(*calls, sig=None):
    """A model turn with function calls: fc(("name", {args}, "id"), ...)."""
    parts = []
    for i, call in enumerate(calls):
        name, args, cid = (tuple(call) + (None, None))[:3]
        part = {"functionCall": {"name": name, "args": args or {}}}
        if cid:
            part["functionCall"]["id"] = cid
        if sig and i == 0:
            part["thoughtSignature"] = sig
        parts.append(part)
    return {"candidates": [{"content": {"role": "model", "parts": parts},
                            "finishReason": "STOP"}]}


def say(text):
    return {"candidates": [{"content": {"role": "model", "parts": [{"text": text}]},
                            "finishReason": "STOP"}]}


def gemini_error(code, status, message, reason=None):
    err = {"code": code, "message": message, "status": status}
    if reason:
        err["details"] = [{"@type": "type.googleapis.com/google.rpc.ErrorInfo",
                           "reason": reason}]
    return (code, {"error": err})


@pytest.fixture
def gemini():
    srv = FakeServer()
    yield srv
    srv.close()


def _agent(box, gemini, **kw):
    import agent
    from google import genai
    from google.genai import types
    client = genai.Client(api_key="test-key", http_options=types.HttpOptions(
        base_url=gemini.base + "/"))                     # no retries in tests
    return agent.CitationAgent(client, box, **kw)


@needs_genai
def test_agent_uses_tools_and_answers(box, gemini):
    gemini.post_script = [
        fc(("list_documents",), sig="c2lnLTE="),
        fc(("add_footnotes", {"name": "essay"}, "call-2"), sig="c2lnLTI="),
        say("Done! 3 footnotes added."),
    ]
    seen = []
    bot = _agent(box, gemini, on_tool=lambda name, args: seen.append((name, args)))
    assert bot.ask("please add footnotes to my essay") == "Done! 3 footnotes added."
    assert seen == [("list_documents", {}), ("add_footnotes", {"name": "essay"})]
    assert (box.output_dir / "My Essay_with_footnotes.docx").is_file()

    first = gemini.posts[0]
    assert first["path"].endswith("/models/gemini-flash-latest:generateContent")
    assert first["headers"].get("x-goog-api-key") == "test-key"
    body = first["body"]
    assert "Never invent" in body["systemInstruction"]["parts"][0]["text"]
    assert [f["name"] for f in body["tools"][0]["functionDeclarations"]] == [
        "list_documents", "analyze_document", "add_footnotes", "check_links"]

    turns = gemini.posts[2]["body"]["contents"]
    assert [t["role"] for t in turns] == ["user", "model", "user", "model", "user"]
    # Gemini 3 requires its thought signatures back, unchanged
    assert turns[1]["parts"][0]["thoughtSignature"] == "c2lnLTE="
    assert turns[3]["parts"][0]["thoughtSignature"] == "c2lnLTI="
    reply = turns[4]["parts"][0]["functionResponse"]
    assert reply["name"] == "add_footnotes" and reply["id"] == "call-2"
    assert reply["response"]["footnotes_added"] == 3
    assert reply["response"]["validation"] == "PASS"
    assert len(bot.history) == 6


@needs_genai
def test_agent_answers_parallel_calls_in_one_turn(box, gemini):
    gemini.post_script = [
        fc(("list_documents",), ("analyze_document", {"name": "Report"})),
        say("ok"),
    ]
    assert _agent(box, gemini).ask("what do I have?") == "ok"
    replies = gemini.posts[1]["body"]["contents"][-1]["parts"]
    assert [p["functionResponse"]["name"] for p in replies] == [
        "list_documents", "analyze_document"]
    assert replies[1]["functionResponse"]["response"]["citations_found"] == 1


@needs_genai
def test_tool_errors_go_back_to_the_model(box, gemini):
    gemini.post_script = [fc(("add_footnotes", {"name": "thesis"})),
                          say("I couldn't find a document called thesis.")]
    bot = _agent(box, gemini)
    assert "couldn't find" in bot.ask("footnotes for my thesis")
    reply = gemini.posts[1]["body"]["contents"][-1]["parts"][0]["functionResponse"]
    assert "no document called 'thesis'" in reply["response"]["error"].lower()


@needs_genai
@pytest.mark.parametrize("error, advice", [
    (gemini_error(429, "RESOURCE_EXHAUSTED", "Resource has been exhausted"), "free-tier limit"),
    (gemini_error(400, "INVALID_ARGUMENT", "API key not valid. Please pass a valid API key.",
                  "API_KEY_INVALID"), "--forget-key"),
    (gemini_error(401, "UNAUTHENTICATED", "Request had invalid authentication credentials. "
                  "Expected OAuth 2 access token, login cookie or other valid authentication "
                  "credential."), "--forget-key"),
    (gemini_error(404, "NOT_FOUND", "models/gemini-x is not found"), "--model"),
    (gemini_error(503, "UNAVAILABLE", "The model is overloaded."), "busy"),
])
def test_gemini_errors_become_advice_and_roll_back(box, gemini, error, advice):
    import agent
    gemini.post_script = [error]
    bot = _agent(box, gemini)
    with pytest.raises(Exception) as exc:
        bot.ask("hello")
    assert advice in agent.friendly_error(exc.value, "gemini-x")
    assert bot.history == []                             # conversation still valid


def test_network_errors_become_advice():
    import agent
    assert "internet" in agent.friendly_error(ConnectionError("down"))
    assert "internet" in agent.friendly_error(TimeoutError())


@needs_genai
def test_quota_message_mentioning_api_key_is_still_a_rate_limit():
    import agent
    from google.genai import errors
    exc = errors.ClientError(429, {"error": {"code": 429, "status": "RESOURCE_EXHAUSTED",
                                             "message": "Quota exceeded for this API key."}})
    assert "free-tier limit" in agent.friendly_error(exc)


def test_plain_removes_markdown():
    import agent
    md = "## Result\n**3 footnotes** added to __Essay__.\n* one\n  * two\n\u2022 three\n- kept"
    assert agent.plain(md) == ("Result\n3 footnotes added to Essay.\n- one\n  - two\n"
                               "- three\n- kept")
    assert agent.plain("2 * 3 = 6, a*b") == "2 * 3 = 6, a*b"


@needs_genai
def test_empty_answer_and_step_limit_roll_back(box, gemini):
    gemini.post_script = [{"candidates": [{"finishReason": "SAFETY"}]}]
    bot = _agent(box, gemini)
    assert "empty answer" in bot.ask("hi") and bot.history == []
    gemini.post_script = [fc(("list_documents",))] * 3
    bot = _agent(box, gemini, max_steps=3)
    assert "too many steps" in bot.ask("loop") and bot.history == []


@needs_genai
def test_main_once_end_to_end(box, gemini, monkeypatch, capsys):
    import agent
    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    monkeypatch.setenv("GOOGLE_GEMINI_BASE_URL", gemini.base + "/")
    gemini.post_script = [fc(("add_footnotes", {"name": "Report"})), say("All done.")]
    code = agent.main(["--once", "footnotes for Report", "--model", "gemini-test",
                       "--input-dir", str(box.input_dir), "--output-dir", str(box.output_dir)])
    out = capsys.readouterr().out
    assert code == 0
    assert "... adding footnotes to \"Report\"" in out and "All done." in out
    assert (box.output_dir / "Report_with_footnotes.docx").is_file()
    assert "/models/gemini-test:generateContent" in gemini.posts[0]["path"]


def test_api_key_sources(tmp_path, monkeypatch, capsys):
    import agent
    key_file = tmp_path / "home" / "gemini_api_key.txt"
    monkeypatch.setattr(agent, "KEY_FILE", key_file)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    assert agent.get_api_key(interactive=False) == (None, None)
    key_file.parent.mkdir()
    key_file.write_text("  saved-key\n", encoding="utf-8")
    assert agent.get_api_key(interactive=False) == ("saved-key", "file")
    monkeypatch.setenv("GEMINI_API_KEY", "env-key")
    assert agent.get_api_key(interactive=False) == ("env-key", "env")
    assert agent.main(["--forget-key"]) == 0 and not key_file.exists()
    assert "Deleted" in capsys.readouterr().out


def test_default_key_location_is_outside_the_project():
    import agent
    assert agent.HERE not in Path(agent.KEY_FILE).resolve().parents


# --------------------------------------------------------------------------
# API key entry and recovery
# --------------------------------------------------------------------------
GOOD_KEY = "GOOD-key-0123456789abcdefXYZ1"
BAD_KEY = "BAD-key-0123456789abcdefXYZ22"


def _models_get(query, headers):
    """GET /v1beta/models/<model>: accepts only GOOD_KEY (like Google)."""
    if headers.get("x-goog-api-key") == GOOD_KEY:
        return (200, JSON, {"name": "models/gemini-flash-latest"})
    return (401, JSON, {"error": {"code": 401, "status": "UNAUTHENTICATED", "message": (
        "Request had invalid authentication credentials. Expected OAuth 2 access "
        "token, login cookie or other valid authentication credential.")}})


def _chat_reply(text):
    """POST generateContent that answers ``text`` only for GOOD_KEY."""
    def reply(headers):
        if headers.get("x-goog-api-key") == GOOD_KEY:
            return say(text)
        return _models_get({}, headers)[::2]          # (401, error-json)
    return reply


class TTYInput(__import__("io").StringIO):
    def isatty(self):
        return True


@pytest.fixture
def keyenv(tmp_path, monkeypatch, gemini):
    """Gemini fake + temp key file, no key in the environment."""
    import agent
    monkeypatch.setattr(agent, "KEY_FILE", tmp_path / "home" / "gemini_api_key.txt")
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    monkeypatch.setenv("GOOGLE_GEMINI_BASE_URL", gemini.base + "/")
    gemini.routes["/v1beta/models/gemini-flash-latest"] = _models_get
    return agent


def _script_getpass(monkeypatch, agent, keys):
    it = iter(keys)
    monkeypatch.setattr(agent.getpass, "getpass", lambda prompt="": next(it))


def test_key_problem_and_mask():
    import agent
    assert agent.key_problem("") == "nothing was pasted"
    assert "didn't work" in agent.key_problem("\x16")          # Ctrl+V as a character
    assert "didn't work" in agent.key_problem("AIza123 456789012345678")
    assert "only 9 characters" in agent.key_problem("123456789")
    assert agent.key_problem(GOOD_KEY) == ""
    assert agent.mask(GOOD_KEY) == "GOOD...XYZ1 (29 characters)"
    assert GOOD_KEY not in agent.mask(GOOD_KEY)
    assert agent.mask("short") == "(5 characters)"


@needs_genai
def test_check_key(keyenv, gemini):
    agent = keyenv
    assert agent.check_key(GOOD_KEY) == (True, "")
    ok, reason = agent.check_key(BAD_KEY)
    assert ok is False and "didn't accept" in reason
    gemini.routes["/v1beta/models/gemini-x"] = (404, JSON, {"error": {
        "code": 404, "status": "NOT_FOUND", "message": "models/gemini-x is not found"}})
    assert agent.check_key(GOOD_KEY, model="gemini-x") == (True, "")   # key is fine
    gemini.routes["/v1beta/models/gemini-y"] = (503, JSON, {"error": {
        "code": 503, "status": "UNAVAILABLE", "message": "overloaded"}})
    ok, reason = agent.check_key(GOOD_KEY, model="gemini-y")
    assert ok is None and "busy" in reason


@needs_genai
def test_bad_pastes_are_rejected_and_only_a_working_key_is_saved(keyenv, monkeypatch, capsys):
    agent = keyenv
    _script_getpass(monkeypatch, agent, ["\x16", BAD_KEY, GOOD_KEY])
    monkeypatch.setattr("builtins.input", lambda prompt="": "y")
    assert agent.get_api_key(interactive=True) == (GOOD_KEY, "typed")
    out = capsys.readouterr().out
    assert "That isn't a key" in out                       # failed Ctrl+V
    assert "Google didn't accept this key" in out          # wrong key, never saved
    assert "Got GOOD...XYZ1 (29 characters)" in out and "The key works." in out
    assert GOOD_KEY not in out and BAD_KEY not in out      # keys are never printed
    assert agent.KEY_FILE.read_text(encoding="utf-8") == GOOD_KEY


@needs_genai
def test_three_wrong_keys_give_up_without_saving(keyenv, monkeypatch, capsys):
    agent = keyenv
    _script_getpass(monkeypatch, agent, [BAD_KEY, "", BAD_KEY])
    assert agent.get_api_key(interactive=True) == (None, None)
    assert "No working key yet" in capsys.readouterr().out
    assert not agent.KEY_FILE.exists()


@needs_genai
def test_rejected_saved_key_is_replaced_in_the_chat(keyenv, gemini, box, monkeypatch, capsys):
    agent = keyenv
    agent.KEY_FILE.parent.mkdir(parents=True)
    agent.KEY_FILE.write_text(BAD_KEY, encoding="utf-8")   # the user's mistake
    gemini.post_script = [_chat_reply("Hi there!"), _chat_reply("Hi there!")]
    _script_getpass(monkeypatch, agent, [GOOD_KEY])
    # "hey" -> 401 -> "Paste a new key now?" y -> paste -> "Save?" y -> retried
    monkeypatch.setattr(sys, "stdin", TTYInput("hey\ny\ny\nexit\n"))
    code = agent.main(["--input-dir", str(box.input_dir), "--output-dir", str(box.output_dir)])
    out = capsys.readouterr().out
    assert code == 0
    assert "didn't accept the API key saved on this computer" in out
    assert "Removed the old saved key." in out and "The key works." in out
    assert "Trying your message again..." in out and "Agent> Hi there!" in out
    assert agent.KEY_FILE.read_text(encoding="utf-8") == GOOD_KEY
    assert [p["headers"]["x-goog-api-key"] for p in gemini.posts] == [BAD_KEY, GOOD_KEY]


@needs_genai
def test_declining_a_new_key_keeps_the_chat_going(keyenv, gemini, box, monkeypatch, capsys):
    agent = keyenv
    agent.KEY_FILE.parent.mkdir(parents=True)
    agent.KEY_FILE.write_text(BAD_KEY, encoding="utf-8")
    gemini.post_script = [_chat_reply("unused")]
    monkeypatch.setattr(sys, "stdin", TTYInput("hey\nn\nexit\n"))
    assert agent.main(["--input-dir", str(box.input_dir),
                       "--output-dir", str(box.output_dir)]) == 0
    out = capsys.readouterr().out
    assert "--forget-key" in out and "Goodbye!" in out
    assert agent.KEY_FILE.read_text(encoding="utf-8") == BAD_KEY   # user said no


@needs_genai
def test_rejected_environment_key_explains_instead_of_prompting(keyenv, gemini, box,
                                                                 monkeypatch, capsys):
    agent = keyenv
    monkeypatch.setenv("GEMINI_API_KEY", BAD_KEY)
    gemini.post_script = [_chat_reply("unused")]
    monkeypatch.setattr(sys, "stdin", TTYInput("hey\nexit\n"))
    assert agent.main(["--input-dir", str(box.input_dir),
                       "--output-dir", str(box.output_dir)]) == 0
    out = capsys.readouterr().out
    assert "environment variable" in out and "Paste a new key now?" not in out


# --------------------------------------------------------------------------
# Other free AIs: Groq, OpenRouter, Mistral (OpenAI format) and Ollama
# --------------------------------------------------------------------------
import llm_providers                                         # noqa: E402
from llm_providers import PROVIDERS, ChatClient, ProviderError  # noqa: E402


def oa_calls(*calls):
    """An OpenAI-style assistant turn with tool calls: oa_calls(("name", {args}, "id"), ...)."""
    tool_calls = []
    for i, call in enumerate(calls):
        name, args, cid = (tuple(call) + (None, None))[:3]
        tool_calls.append({"id": cid or f"call_{i}", "type": "function",
                           "function": {"name": name, "arguments": json.dumps(args or {})}})
    return {"id": "chatcmpl-1", "object": "chat.completion", "model": "m", "choices": [
        {"index": 0, "finish_reason": "tool_calls",
         "message": {"role": "assistant", "content": None, "tool_calls": tool_calls}}]}


def oa_say(text):
    return {"choices": [{"index": 0, "finish_reason": "stop",
                         "message": {"role": "assistant", "content": text}}]}


def oa_error(status, message, code=None, headers=None):
    err = {"message": message, "type": "invalid_request_error"}
    if code:
        err["code"] = code
    return (status, {"error": err}, headers or {})


def ol_calls(*calls):
    """An Ollama /api/chat answer with tool calls (arguments are objects, no ids)."""
    return {"model": "qwen3:8b", "done": True, "done_reason": "stop", "message": {
        "role": "assistant", "content": "",
        "tool_calls": [{"function": {"index": i, "name": name, "arguments": args}}
                       for i, (name, args) in enumerate(calls)]}}


def ol_say(text):
    return {"model": "qwen3:8b", "done": True, "done_reason": "stop",
            "message": {"role": "assistant", "content": text}}


ROOT_PATH = {"gemini": "", "groq": "/openai/v1", "openrouter": "/api/v1",
             "mistral": "/v1", "ollama": ""}
DAILY_429 = ("Rate limit reached for model `openai/gpt-oss-120b` in organization `org_1` "
             "service tier `on_demand` on requests per day (RPD): Limit 1000, Used 1000, "
             "Requested 1. Please try again in 1m26.4s.")


@pytest.fixture
def llm():
    srv = FakeServer()
    yield srv
    srv.close()


@pytest.fixture
def clean(tmp_path, monkeypatch):
    """No real keys, models or addresses from the developer's environment;
    keys are saved in a temp folder."""
    import agent
    monkeypatch.setattr(agent, "KEY_FILE", tmp_path / "home" / "gemini_api_key.txt")
    for p in PROVIDERS.values():
        for var in p.key_envs + (p.model_env, p.base_url_env):
            if var:
                monkeypatch.delenv(var, raising=False)
    monkeypatch.delenv("GOOGLE_GEMINI_BASE_URL", raising=False)
    return agent


def _point(monkeypatch, provider, srv):
    """Send ``provider``'s requests to the fake server."""
    monkeypatch.setenv(PROVIDERS[provider].base_url_env, srv.base + ROOT_PATH[provider])


def _chat_agent(box, srv, provider="groq", key="test-key", sleep=None, on_wait=None, **kw):
    import agent
    p = PROVIDERS[provider]
    client = ChatClient(p, key, base_url=srv.base + ROOT_PATH[provider],
                        sleep=sleep or (lambda s: None), on_wait=on_wait)
    return agent.ChatAgent(client, box, model=kw.pop("model", p.default_model), **kw)


def test_chat_agent_uses_tools_and_answers(box, llm):
    llm.post_script = [oa_calls(("list_documents", {}, "c1")),
                       oa_calls(("add_footnotes", {"name": "essay"}, "c2")),
                       oa_say("Done! 3 footnotes added.")]
    seen = []
    bot = _chat_agent(box, llm, on_tool=lambda name, args: seen.append((name, args)))
    assert bot.ask("please add footnotes to my essay") == "Done! 3 footnotes added."
    assert seen == [("list_documents", {}), ("add_footnotes", {"name": "essay"})]
    assert (box.output_dir / "My Essay_with_footnotes.docx").is_file()

    first = llm.posts[0]
    assert first["path"] == "/openai/v1/chat/completions"
    assert first["headers"]["Authorization"] == "Bearer test-key"
    body = first["body"]
    assert body["model"] == "openai/gpt-oss-120b" and body["reasoning_effort"] == "low"
    assert body["max_completion_tokens"] == llm_providers.GROQ_MAX_OUTPUT
    assert body["messages"][0]["role"] == "system"
    assert "Never invent" in body["messages"][0]["content"]
    assert [t["function"]["name"] for t in body["tools"]] == [
        "list_documents", "analyze_document", "add_footnotes", "check_links"]
    assert body["tools"][0]["function"]["parameters"] == {"type": "object", "properties": {}}

    msgs = llm.posts[2]["body"]["messages"]
    assert [m["role"] for m in msgs] == ["system", "user", "assistant", "tool",
                                         "assistant", "tool"]
    assert msgs[4]["tool_calls"][0]["id"] == "c2" and msgs[4]["content"] is None
    assert msgs[5]["tool_call_id"] == "c2" and msgs[5]["name"] == "add_footnotes"
    result = json.loads(msgs[5]["content"])
    assert result["footnotes_added"] == 3 and result["validation"] == "PASS"
    assert len(bot.history) == 6


def test_chat_agent_parallel_calls_bad_arguments_and_thinking(box, llm):
    llm.post_script = [
        {"choices": [{"finish_reason": "tool_calls", "message": {
            "role": "assistant", "content": "", "tool_calls": [
                {"id": "a", "type": "function",
                 "function": {"name": "list_documents", "arguments": "null"}},
                {"id": "b", "type": "function",
                 "function": {"name": "analyze_document", "arguments": "{\"name\": \"Report\"}"}},
                {"id": "c", "type": "function",
                 "function": {"name": "analyze_document", "arguments": "{not json"}}]}}]},
        oa_say("<think>which one?</think>ok"),
    ]
    bot = _chat_agent(box, llm, provider="openrouter")
    assert bot.ask("what do I have?") == "ok"               # thinking is not shown
    assert llm.posts[0]["path"] == "/api/v1/chat/completions"
    assert "reasoning_effort" not in llm.posts[0]["body"]
    assert "max_completion_tokens" not in llm.posts[0]["body"]
    tools = [m for m in llm.posts[1]["body"]["messages"] if m["role"] == "tool"]
    assert [m["tool_call_id"] for m in tools] == ["a", "b", "c"]
    assert json.loads(tools[0]["content"])["documents"] == ["My Essay.docx", "Report.docx"]
    assert json.loads(tools[1]["content"])["citations_found"] == 1
    assert "not a valid JSON" in json.loads(tools[2]["content"])["error"]


def test_ollama_uses_its_own_api_with_a_big_enough_context(box, llm):
    llm.post_script = [ol_calls(("add_footnotes", {"name": "Report"})), ol_say("All done.")]
    bot = _chat_agent(box, llm, provider="ollama", key="")
    assert bot.ask("footnotes for Report") == "All done."
    first = llm.posts[0]
    assert first["path"] == "/api/chat" and "Authorization" not in first["headers"]
    assert first["body"]["stream"] is False
    assert first["body"]["options"]["num_ctx"] == llm_providers.OLLAMA_NUM_CTX
    msgs = llm.posts[1]["body"]["messages"]
    assert msgs[2]["tool_calls"][0]["function"]["arguments"] == {"name": "Report"}
    assert msgs[3]["role"] == "tool" and msgs[3]["tool_name"] == "add_footnotes"
    assert json.loads(msgs[3]["content"])["validation"] == "PASS"
    assert bot.history[1]["tool_calls"][0]["id"]             # an id was made up


def test_per_minute_limit_is_waited_for_then_retried(box, llm):
    waits, slept = [], []
    llm.post_script = [
        oa_error(429, "Rate limit reached ... on tokens per minute (TPM): Limit 8000. "
                 "Please try again in 7.66s.", "rate_limit_exceeded", {"retry-after": "8"}),
        oa_error(429, "Rate limit reached ... Please try again in 2.5s.", "rate_limit_exceeded"),
        oa_say("hi"),
    ]
    bot = _chat_agent(box, llm, sleep=slept.append,
                      on_wait=lambda p, seconds, e: waits.append((p.id, seconds)))
    assert bot.ask("hello") == "hi"
    assert waits == [("groq", 8.5), ("groq", 3.0)] and slept == [8.5, 3.0]


@pytest.mark.parametrize("provider, error, advice", [
    ("groq", oa_error(401, "Invalid API Key", "invalid_api_key"),
     "python agent.py --provider groq --forget-key"),
    ("groq", oa_error(429, DAILY_429, "rate_limit_exceeded"), "free requests for today"),
    ("groq", oa_error(404, "The model `x` does not exist or you do not have access to it.",
                      "model_not_found"), "--provider groq --model openai/gpt-oss-120b"),
    ("groq", oa_error(413, "Request too large for model `openai/gpt-oss-120b` on tokens per "
                      "minute (TPM): Limit 8000, Requested 9500"), "too much text"),
    ("openrouter", (404, {"error": {"code": 404, "message": (
        "No endpoints found matching your data policy (Free model training). Configure: "
        "https://openrouter.ai/settings/privacy")}}), "openrouter.ai/settings/privacy"),
    ("openrouter", (404, {"error": {"code": 404, "message":
                                    "No endpoints found that support tool use."}}),
     "can't use tools"),
    ("openrouter", (402, {"error": {"code": 402, "message": "Insufficient credits"}}),
     "no credit left"),
    ("mistral", (400, {"object": "error", "message": "Invalid model: mistral-x",
                       "type": "invalid_model"}), "--provider mistral --model mistral-small"),
    ("mistral", (401, {"message": "Unauthorized", "request_id": "r1"}), "--forget-key"),
    ("ollama", (404, {"error": "model \"qwen3:8b\" not found, try pulling it first"}),
     "ollama pull qwen3:8b"),
    ("ollama", (400, {"error": "registry.ollama.ai/library/gemma3:4b does not support tools"}),
     "can't use tools"),
])
def test_provider_errors_become_advice_and_roll_back(box, llm, provider, error, advice):
    import agent
    llm.post_script = [error]
    bot = _chat_agent(box, llm, provider=provider)
    with pytest.raises(ProviderError) as exc:
        bot.ask("hello")
    assert advice in agent.friendly_error(exc.value, bot.model, provider)
    assert bot.history == []                               # conversation still valid
    assert len(llm.posts) == 1                             # not retried
    assert agent.is_key_error(exc.value) == (error[0] == 401)


def test_busy_service_is_retried_then_explained(box, llm):
    import agent
    llm.post_script = [(503, {"message": "Service unavailable"})] * 4
    slept = []
    bot = _chat_agent(box, llm, provider="mistral", sleep=slept.append)
    with pytest.raises(ProviderError) as exc:
        bot.ask("hello")
    assert slept == [2, 5, 10] and len(llm.posts) == 4
    assert "busy" in agent.friendly_error(exc.value, bot.model, "mistral")
    assert agent.is_limit_error(exc.value)


def test_network_problems_are_explained(box):
    import agent
    port = _closed_port()
    for provider, advice in [("groq", "internet"), ("ollama", "Ollama isn't running")]:
        client = ChatClient(PROVIDERS[provider], "k", base_url=f"http://127.0.0.1:{port}",
                            sleep=lambda s: None)
        bot = agent.ChatAgent(client, box, model="m")
        with pytest.raises(ProviderError) as exc:
            bot.ask("hello")
        assert exc.value.network and advice in agent.friendly_error(exc.value, "m", provider)


def test_error_details_are_understood(monkeypatch):
    assert llm_providers._seconds("1m26.4s") == pytest.approx(86.4)
    assert llm_providers._seconds("500ms") == pytest.approx(0.5)
    e = llm_providers.error_from_http("groq", 429, json.dumps(
        {"error": {"message": "Please try again in 1m26.4s."}}).encode())
    assert e.retry_after == pytest.approx(86.4) and llm_providers.is_daily_limit(e)
    e = llm_providers.error_from_http("openrouter", 429, json.dumps({"error": {
        "code": 429, "message": "Provider returned error",
        "metadata": {"raw": "upstream is rate-limited"}}}).encode())
    assert "upstream is rate-limited" in e.message and not llm_providers.is_daily_limit(e)
    monkeypatch.setenv("OLLAMA_HOST", "0.0.0.0")
    assert llm_providers.base_url_for(PROVIDERS["ollama"]) == "http://127.0.0.1:11434"
    monkeypatch.setenv("OLLAMA_HOST", "otherpc:12345")
    assert llm_providers.base_url_for(PROVIDERS["ollama"]) == "http://otherpc:12345"


def test_errors_reported_with_status_200_are_errors(box, llm):
    llm.post_script = [{"error": {"code": 502, "message": "Upstream error"}}] * 4
    with pytest.raises(ProviderError) as exc:
        _chat_agent(box, llm, provider="openrouter").ask("hello")
    assert exc.value.status == 502


def _msgs_for_exchange(n, big):
    return [{"role": "user", "content": f"q{n}"},
            {"role": "assistant", "content": "", "tool_calls": [{
                "id": f"t{n}", "type": "function",
                "function": {"name": "list_documents", "arguments": "{}"}}]},
            {"role": "tool", "tool_call_id": f"t{n}", "name": "list_documents", "content": big},
            {"role": "assistant", "content": f"a{n}"}]


def test_long_chats_are_trimmed_for_small_free_tiers(box, llm):
    import agent
    big = json.dumps({"data": "x" * 3000})
    bot = _chat_agent(box, llm)
    bot.history = (_msgs_for_exchange(1, big) + _msgs_for_exchange(2, big)
                   + _msgs_for_exchange(3, big)[:3])        # q3 is still being answered
    saved = json.dumps(bot.history)
    base = len(agent.SYSTEM_PROMPT) + len(json.dumps(bot.tools))
    total = base + sum(len(json.dumps(m)) for m in bot.history)

    bot.max_chars = total - 1000                    # one old result has to go
    w = bot.window()
    assert w[2]["content"] == agent.LEFT_OUT and w[6]["content"] == big
    assert w[-1]["content"] == big                  # the current request is never trimmed
    assert [m["role"] for m in w] == [m["role"] for m in bot.history]

    def size(msgs):
        return sum(len(json.dumps(m)) for m in msgs)

    stubbed = [dict(m, content=agent.LEFT_OUT) if m["role"] == "tool" else m
               for m in bot.history[:8]] + bot.history[8:]
    first_exchange = size(stubbed[:4])
    bot.max_chars = base + size(stubbed) - first_exchange // 2   # trimming is not enough
    w = bot.window()
    assert w[0] == {"role": "user", "content": "q2"}             # oldest exchange left out
    assert w[2]["content"] == agent.LEFT_OUT and w[-1]["content"] == big

    bot.max_chars = 1                               # only the current request is left
    assert [m.get("content") for m in bot.window()][0] == "q3"
    assert json.dumps(bot.history) == saved         # the history itself is kept

    bot.history = _msgs_for_exchange(1, big)        # end to end: what is really sent
    llm.post_script = [oa_say("fine")]
    assert bot.ask("q2") == "fine"
    sent = llm.posts[-1]["body"]["messages"]
    assert [(m["role"], m["content"]) for m in sent[1:]] == [("user", "q2")]
    assert len(bot.history) == 6                    # nothing is forgotten locally

    assert _chat_agent(box, llm).max_chars == PROVIDERS["groq"].max_chars == 16_000
    assert _chat_agent(box, llm, provider="ollama").max_chars == 40_000


def test_check_key_for_other_providers(clean, llm, monkeypatch):
    agent = clean

    def accepts_good(query, headers):
        if headers.get("Authorization") == f"Bearer {GOOD_KEY}":
            return (200, JSON, {"data": []})
        return (401, JSON, {"error": {"message": "Invalid API Key",
                                      "code": "invalid_api_key"}})

    llm.routes["/openai/v1/models"] = accepts_good
    llm.routes["/api/v1/key"] = accepts_good
    _point(monkeypatch, "groq", llm)
    _point(monkeypatch, "openrouter", llm)
    assert agent.check_key(GOOD_KEY, provider="groq") == (True, "")
    assert agent.check_key(BAD_KEY, provider="groq") == (False, "Groq didn't accept this key.")
    assert agent.check_key(BAD_KEY, provider="openrouter") == (
        False, "OpenRouter didn't accept this key.")
    monkeypatch.setenv("MISTRAL_BASE_URL", f"http://127.0.0.1:{_closed_port()}/v1")
    ok, reason = agent.check_key(GOOD_KEY, provider="mistral")
    assert ok is None and "internet" in reason


def test_keys_are_saved_per_provider(clean, llm, monkeypatch, capsys):
    agent = clean
    llm.routes["/openai/v1/models"] = lambda q, h: (
        (200, JSON, {"data": []}) if h.get("Authorization") == f"Bearer {GOOD_KEY}"
        else (401, JSON, {"error": {"message": "Invalid API Key"}}))
    _point(monkeypatch, "groq", llm)
    agent.KEY_FILE.parent.mkdir(parents=True)
    agent.KEY_FILE.write_text("gemini-key-0123456789", encoding="utf-8")
    _script_getpass(monkeypatch, agent, [BAD_KEY, GOOD_KEY])
    monkeypatch.setattr("builtins.input", lambda prompt="": "y")
    assert agent.get_api_key(interactive=True, provider="groq") == (GOOD_KEY, "typed")
    out = capsys.readouterr().out
    assert "https://console.groq.com/keys" in out and "doesn't train" in out
    assert "Groq didn't accept this key" in out and "Checking it with Groq" in out
    groq_file = agent.KEY_FILE.parent / "groq_api_key.txt"
    assert groq_file.read_text(encoding="utf-8") == GOOD_KEY
    assert agent.get_api_key(interactive=False, provider="groq") == (GOOD_KEY, "file")
    monkeypatch.setenv("GROQ_API_KEY", "from-env")
    assert agent.get_api_key(interactive=False, provider="groq") == ("from-env", "env")

    assert agent.main(["--provider", "groq", "--forget-key"]) == 0
    assert not groq_file.exists() and agent.KEY_FILE.exists()   # Gemini's key stays
    assert agent.main(["--provider", "ollama", "--forget-key"]) == 0
    assert "doesn't use an API key" in capsys.readouterr().out


def test_main_once_with_groq(clean, llm, box, monkeypatch, capsys):
    agent = clean
    monkeypatch.setenv("GROQ_API_KEY", "groq-key")
    _point(monkeypatch, "groq", llm)
    llm.post_script = [oa_calls(("add_footnotes", {"name": "Report"})), oa_say("**All** done.")]
    code = agent.main(["--provider", "groq", "--once", "footnotes for Report",
                       "--input-dir", str(box.input_dir), "--output-dir", str(box.output_dir)])
    out = capsys.readouterr().out
    assert code == 0 and "All done." in out and "**" not in out
    assert "... adding footnotes to \"Report\"" in out
    assert (box.output_dir / "Report_with_footnotes.docx").is_file()
    assert llm.posts[0]["headers"]["Authorization"] == "Bearer groq-key"


def test_ollama_must_be_running_and_have_the_model(clean, llm, box, monkeypatch, capsys):
    agent = clean
    dirs = ["--input-dir", str(box.input_dir), "--output-dir", str(box.output_dir)]
    monkeypatch.setenv("OLLAMA_HOST", f"127.0.0.1:{_closed_port()}")
    assert agent.main(["--provider", "ollama", "--once", "hi"] + dirs) == 2
    assert "Ollama isn't running" in capsys.readouterr().out

    _point(monkeypatch, "ollama", llm)
    llm.routes["/api/tags"] = (200, JSON, {"models": [{"name": "granite4.1:3b"}]})
    assert agent.main(["--provider", "ollama", "--once", "hi"] + dirs) == 2
    out = capsys.readouterr().out
    assert "ollama pull qwen3:8b" in out and "granite4.1:3b" in out

    llm.post_script = [ol_say("Hello from your computer.")]
    assert agent.main(["--provider", "ollama", "--model", "granite4.1:3b",
                       "--once", "hi"] + dirs) == 0
    assert "Hello from your computer." in capsys.readouterr().out
    assert llm.posts[0]["body"]["model"] == "granite4.1:3b"


def test_use_command_switches_ai_and_keeps_the_conversation(clean, box, monkeypatch, capsys):
    agent = clean
    groq, ollama = FakeServer(), FakeServer()
    try:
        monkeypatch.setenv("GROQ_API_KEY", "groq-key")
        _point(monkeypatch, "groq", groq)
        _point(monkeypatch, "ollama", ollama)
        groq.post_script = [oa_say("Hi! I can add footnotes.")]
        ollama.routes["/api/tags"] = (200, JSON, {"models": [{"name": "qwen3:8b"}]})
        ollama.post_script = [ol_say("You said hello.")]
        monkeypatch.setattr(sys, "stdin", TTYInput(
            "hello\n/use nothing\n/use ollama\nwhat did I say?\nexit\n"))
        assert agent.main(["--provider", "groq", "--input-dir", str(box.input_dir),
                           "--output-dir", str(box.output_dir)]) == 0
        out = capsys.readouterr().out
        assert "Brain: Groq, model openai/gpt-oss-120b" in out and "/use ollama" in out
        assert "Type /use and one of" in out
        assert "Now using Ollama (on your computer), model qwen3:8b" in out
        assert "Agent> You said hello." in out
        msgs = ollama.posts[0]["body"]["messages"]
        assert [(m["role"], m["content"]) for m in msgs[1:]] == [
            ("user", "hello"), ("assistant", "Hi! I can add footnotes."),
            ("user", "what did I say?")]
    finally:
        groq.close()
        ollama.close()


def test_used_up_limit_offers_an_ai_you_have_a_key_for(clean, box, monkeypatch, capsys):
    agent = clean
    groq, mistral = FakeServer(), FakeServer()
    try:
        monkeypatch.setenv("GROQ_API_KEY", "groq-key")
        monkeypatch.setenv("MISTRAL_API_KEY", "mistral-key")
        _point(monkeypatch, "groq", groq)
        _point(monkeypatch, "mistral", mistral)
        groq.post_script = [oa_error(429, DAILY_429, "rate_limit_exceeded")]
        mistral.post_script = [oa_say("Hello from Mistral.")]
        monkeypatch.setattr(sys, "stdin", TTYInput("hello\ny\nexit\n"))
        assert agent.main(["--provider", "groq", "--input-dir", str(box.input_dir),
                           "--output-dir", str(box.output_dir)]) == 0
        out = capsys.readouterr().out
        assert "used up Groq's free requests for today" in out
        assert "Now using Mistral, model mistral-small-latest" in out
        assert "Trying your message again..." in out and "Agent> Hello from Mistral." in out
        assert len(groq.posts) == 1                        # a daily limit isn't retried
        assert mistral.posts[0]["headers"]["Authorization"] == "Bearer mistral-key"
    finally:
        groq.close()
        mistral.close()


def test_used_up_limit_without_other_keys_gives_a_tip(clean, box, llm, monkeypatch, capsys):
    agent = clean
    monkeypatch.setenv("GROQ_API_KEY", "groq-key")
    _point(monkeypatch, "groq", llm)
    llm.post_script = [oa_error(429, DAILY_429, "rate_limit_exceeded")]
    monkeypatch.setattr(sys, "stdin", TTYInput("hello\nexit\n"))
    assert agent.main(["--provider", "groq", "--input-dir", str(box.input_dir),
                       "--output-dir", str(box.output_dir)]) == 0
    out = capsys.readouterr().out
    assert "Tip: you can carry on with another free AI" in out
    assert "/use gemini, /use openrouter, /use mistral, /use ollama" in out


def test_providers_list(clean, monkeypatch, capsys):
    agent = clean
    agent.KEY_FILE.parent.mkdir(parents=True)
    agent.KEY_FILE.write_text("saved-gemini-key-012345", encoding="utf-8")
    monkeypatch.setenv("MISTRAL_API_KEY", "m")
    assert agent.main(["--providers"]) == 0
    out = capsys.readouterr().out
    for pid in PROVIDERS:
        assert f"  {pid} - " in out
    assert "gemini - Google Gemini, the default (key saved)" in out
    assert "groq - Groq (no key yet)" in out and "https://console.groq.com/keys" in out
    assert "mistral - Mistral (key in an environment variable)" in out
    assert "(no key needed)" in out and "Nothing leaves your computer." in out


@needs_genai
def test_gemini_continues_a_conversation_from_another_ai(box, gemini):
    gemini.post_script = [say("You said hello.")]
    bot = _agent(box, gemini)
    bot.seed([("hello", "Hi! I can add footnotes.")])
    assert bot.ask("what did I say?") == "You said hello."
    turns = gemini.posts[0]["body"]["contents"]
    assert [(t["role"], t["parts"][0]["text"]) for t in turns] == [
        ("user", "hello"), ("model", "Hi! I can add footnotes."), ("user", "what did I say?")]


def test_waiting_is_shown(capsys):
    import agent
    agent._show_wait(PROVIDERS["groq"], 7.2, ProviderError("groq", 429, "slow down"))
    agent._show_wait(PROVIDERS["mistral"], 2, ProviderError("mistral", 503, "busy"))
    out = capsys.readouterr().out
    assert "  ... Groq's per-minute limit is reached - waiting 8 seconds" in out
    assert "  ... Mistral is busy - waiting 2 seconds" in out


def test_chat_agent_short_empty_and_endless_answers(box, llm):
    llm.post_script = [{"choices": [{"finish_reason": "length", "message": {
        "role": "assistant", "content": "Problems: [5] is not in"}}]}]
    bot = _chat_agent(box, llm)
    answer = bot.ask("list the problems")
    assert answer.startswith("Problems: [5] is not in") and "cut short" in answer
    assert bot.history[-1]["content"] == "Problems: [5] is not in"   # the note stays local

    llm.post_script = [{"choices": [{"finish_reason": "content_filter", "message": {
        "role": "assistant", "content": ""}}]}]
    bot = _chat_agent(box, llm, provider="mistral")
    assert bot.ask("hi") == ("(Mistral sent an empty answer - reason: content_filter. "
                             "Please try asking in a different way.)")
    assert bot.history == []

    llm.post_script = [oa_calls(("list_documents",))] * 3
    bot = _chat_agent(box, llm, max_steps=3)
    assert "too many steps" in bot.ask("loop") and bot.history == []
