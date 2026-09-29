"""Validation of the generated output: real footnotes, document preservation,
marker replacement, existing-footnote preservation, and no fabrication.

Run AFTER `python tests/create_sample.py` and `python main.py`.
Exits 0 when every check passes, 1 otherwise.
"""
import re
import sys
import zipfile
from lxml import etree

from docx import Document
from docx.oxml.ns import qn

OUT = "output/sample_with_footnotes.docx"
IN = "input/sample.docx"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# Expected: 11 citations in the body (range [4]-[6] expands to 3) -> 11 new
# footnotes; plus 1 pre-existing manual footnote = 12 regular footnotes and
# 12 body footnoteReference elements.
EXPECTED_REGULAR = 12
EXPECTED_BODY_REFS = 12
EXPECTED_NEW_URLS = 9   # refs 1-9 (ref 10 has no URL; ref 11 is uncited)
PREEXISTING_TEXT = "This footnote already existed in the source document."


def tag(e):
    return e.tag


def regular_footnotes(path):
    z = zipfile.ZipFile(path)
    name = [n for n in z.namelist() if n.endswith("footnotes.xml")]
    if not name:
        return None
    return etree.fromstring(z.read(name[0]))


def fn_defs(root):
    return [f for f in root.findall(qn("w:footnote"))
            if f.get(qn("w:id")) not in ("-1", "0")]


def fn_text(f):
    return "".join(t.text or "" for t in f.iter(qn("w:t")))


def body_ref_count(doc):
    n = 0
    for p in doc.paragraphs:
        n += len(p._p.findall(".//" + qn("w:footnoteReference")))
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    n += len(p._p.findall(".//" + qn("w:footnoteReference")))
    return n


def main():
    ok = True

    def check(name, cond, detail=""):
        nonlocal ok
        print(f"[{'PASS' if cond else 'FAIL'}] {name}{(' - ' + detail) if detail else ''}")
        ok &= bool(cond)

    # 1) Input still contains ONLY its one pre-existing footnote ------------
    zin_ft = regular_footnotes(IN)
    check("input has footnotes.xml", zin_ft is not None)
    in_defs = {f.get(qn("w:id")): fn_text(f) for f in fn_defs(zin_ft)}
    check("input has exactly 1 regular footnote (the manual one)",
          in_defs == {"1": PREEXISTING_TEXT}, str(in_defs))
    in_doc = Document(IN)
    check("input body has exactly 1 footnoteReference",
          body_ref_count(in_doc) == 1, str(body_ref_count(in_doc)))

    # 2) Output validity + real footnotes ------------------------------------
    check("output is a valid zip", zipfile.is_zipfile(OUT))
    doc = Document(OUT)  # must not raise
    check("output reopens with python-docx", True)
    ft = regular_footnotes(OUT)
    check("word/footnotes.xml present in output", ft is not None)
    regular = fn_defs(ft)
    check("regular footnotes == 12", len(regular) == EXPECTED_REGULAR,
          str(len(regular)))
    check("body footnoteReference == 12",
          body_ref_count(doc) == EXPECTED_BODY_REFS, str(body_ref_count(doc)))

    # 3) No fabricated URLs: every new footnote URL comes from a reference ---
    ref_urls = {
        "https://example.com/one", "https://example.com/two",
        "https://doi.org/10.1000/three", "https://example.com/four",
        "https://example.com/five", "https://example.com/six",
        "https://example.com/seven", "https://example.com/eight",
        "https://example.com/nine",
    }
    new_texts = [fn_text(f) for f in regular
                 if f.get(qn("w:id")) != "1"]
    check("11 new footnotes", len(new_texts) == 11, str(len(new_texts)))
    check("all new footnote URLs are from the reference list",
          set(new_texts) <= ref_urls, str(sorted(set(new_texts) - ref_urls)))
    check("9 distinct new URLs", len(set(new_texts)) == EXPECTED_NEW_URLS,
          str(len(set(new_texts))))

    # 4) Pre-existing footnote preserved verbatim ----------------------------
    defs = {f.get(qn("w:id")): fn_text(f) for f in regular}
    check("pre-existing footnote (id 1) preserved",
          defs.get("1") == PREEXISTING_TEXT, repr(defs.get("1")))

    # 5) No invalid style usage (a paragraph style must never be an rStyle) --
    bad = []
    for f in regular:
        for rs in f.iter(qn("w:rStyle")):
            if (rs.get(qn("w:val")) or "").lower() == "footnotetext":
                bad.append(rs.get(qn("w:val")))
    check("no w:rStyle pointing at the FootnoteText paragraph style", not bad, str(bad))

    # 6) Hyperlink preserved ---------------------------------------------------
    hl = doc.element.body.findall(".//" + qn("w:hyperlink"))
    check("hyperlink element preserved", len(hl) >= 1, str(len(hl)))
    targets = [doc.part.rels[r.get(qn("r:id"))].target_ref
               for r in hl if r.get(qn("r:id")) in doc.part.rels]
    check("hyperlink target preserved",
          "https://example.com/external" in targets, str(targets))

    # 7) Bold formatting preserved around the [7] citation -------------------
    bold_ok = False
    for p in doc.paragraphs:
        for r in p.runs:
            if r.bold and "inside it" in (r.text or ""):
                bold_ok = True
    check("bold run 'inside it.' preserved", bold_ok)

    # 8) Table-cell citation got a footnote -----------------------------------
    table_fn = 0
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    table_fn += len(p._p.findall(".//" + qn("w:footnoteReference")))
    check("table-cell footnote references == 1", table_fn == 1, str(table_fn))

    # 9) Markers replaced in the body (before the References heading) ---------
    body_text = "\n".join(p.text for p in doc.paragraphs)
    idx = body_text.find("References")
    leftover = re.findall(r"\[\d+\]", body_text[:idx])
    check("no [n] markers left in body (replace mode)", not leftover, str(leftover))

    print("\nALL CHECKS PASSED" if ok else "\nSOME CHECKS FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
